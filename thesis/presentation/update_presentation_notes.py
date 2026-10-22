import os
import pptx

SPEECH_NOTES_SIMPLE_ENGLISH = [
    # SLIDE 1: Title & Welcome
    """[SLIDE 1 | TIME: 0:00 - 0:30 | TARGET: 30 SECONDS]

SPEECH (EN):
"Hello everyone, my name is Alessio Curcio, and today I am pleased to present my internship project titled:
'Seismic Viewer: Design, Implementation and Evaluation of an Interactive 4D Spatio-Temporal Visualisation System for Seismological Data'.

This work was carried out during my curricular internship at the National Institute of Geophysics and Volcanology (INGV), under the supervision of Prof. Danilo Avola from Sapienza and Dr. Daniele Bailo from INGV."

[VISUAL ACTION: Stand upright, smile, make warm eye contact with the Committee, project confidence]""",

    # SLIDE 2: Domain Context & Research Motivation
    """[SLIDE 2 | TIME: 0:30 - 1:15 | TARGET: 45 SECONDS]

SPEECH (EN):
"To understand the motivation of this project, we look at the ongoing volcanic unrest at Campi Flegrei near Naples.
Since 2005, continuous ground uplift accompanied by dense earthquake swarms has required multi-parametric monitoring.

Recently, in 2025, a landmark study published in Nature Communications Earth & Environment by Dr. Guido Giordano and Dr. Lucilla Alfonsi showed that these earthquakes aren't just diffuse clusters: they outline active, growing planar fault structures.

During my internship, I worked in direct collaboration with Giordano and Alfonsi.
Our goal was to solve an operational need at INGV: building an interactive desktop system that bridges raw earthquake catalogues, subsurface tomography, and automated fault plane extraction in one responsive workspace."

[VISUAL ACTION: Point to Active Volcanic Unrest on Card 1, highlight the Giordano & Alfonsi Nature 2025 study on Card 2, and point to INGV Operational Need on Card 3]""",

    # SLIDE 3: State of the Art & Tooling Comparison
    """[SLIDE 3 | TIME: 1:15 - 1:55 | TARGET: 40 SECONDS]

SPEECH (EN):
"When we analyzed existing seismological tools, we found a clear functional gap:

First, General GIS platforms like QGIS and ArcGIS have a 2D projection bias that flattens depth into static plan-views. When navigating dense 3D point clouds over 10,000 events, they suffer major frame drops and lack native 4D time playback.
Second, Scientific Code like ObsPy and GMT produces static image files in batch mode without real-time interaction, requiring continuous scripting. Even 3D tools like ParaView involve complex manual pipelines and lack direct seismological integrations.

Seismic Viewer bridges this gap: providing hardware-accelerated 2D, 3D, and 4D WebGL navigation, automated fault plane fitting, and an easy standalone application without heavy GIS overhead."

[VISUAL ACTION: Walk through the three columns on Slide 3: GIS platforms, Scientific Code (ObsPy/GMT/ParaView), and our solution Seismic Viewer]""",

    # SLIDE 4: System Design & Implementation Timeline
    """[SLIDE 4 | TIME: 1:55 - 2:40 | TARGET: 45 SECONDS]

SPEECH (EN):
"We defined three core functional pillars:
Hardware-accelerated Interactive Visualisations in 2D, 3D, and 4D;
Automated Structural Analysis combining spatial clustering with RANSAC-SVD plane fitting;
and Multi-Layer Geophysical Fusion combining tomography, ground uplift, and seismic stations.

Working closely with researchers Giordano and Alfonsi, development progressed through six iterative phases:
starting with 3D prototypes and CSV parsing in Phase 1,
testing K-Means and the Python-WebView bridge in Phase 2,
building the analytical dashboard and querying the INGV GOSSIP API in Phase 3,
introducing DBSCAN with knee tuning, 4D cylinders, and cluster locking in Phase 4,
integrating Vp/Vs tomography and GNSS surface deformation in Phase 5,
and finalizing high-speed Float32 binary streaming in Phase 6."

[VISUAL ACTION: Point to the three functional pillars on Card 1, then follow the 6-Phase Timeline from Phase 1 to Phase 6 on Card 2]""",

    # SLIDE 5: 2D Application Workspace
    """[SLIDE 5 | TIME: 2:40 - 3:20 | TARGET: 40 SECONDS]

SPEECH (EN):
"Starting with the 2D Application Workspace:

All event filtering across Magnitude, Depth, and Date ranges runs entirely in JavaScript inside the browser. This eliminates HTTP roundtrips during parameter adjustments, providing instant, zero-latency filtering without re-fetching raw catalogues.
For spatial representation, we provide a bi-modal display:
Users can toggle between Proportional Scatter Markers—where marker radius scales with earthquake magnitude—and a Continuous Kernel Density heatmap that aggregates hypocentres into Gaussian density fields.
This allows geophysicists to immediately highlight dense swarm concentrations across Pozzuoli and the Solfatara crater."

[VISUAL ACTION: Point to the interactive drawer filters on the left and show the switch between Proportional Scatter Markers and Continuous Kernel Density]""",

    # SLIDE 6: 3D Application Workspace
    """[SLIDE 6 | TIME: 3:20 - 4:00 | TARGET: 40 SECONDS]

SPEECH (EN):
"Switching to the 3D Workspace, hypocentres are rendered in physical space via WebGL.
To guarantee scientific fidelity, we implemented three key mechanisms:

First, Geographic Coordinate Retention: horizontal axes preserve standard geographic degrees alongside physical focal depth measured in kilometres.
Second, Aspect Ratio Correction: to prevent distortion caused by Earth's sphericity, the system scales the 3D bounding box so that 1 kilometre along X, Y, and Z renders with identical physical proportions.
And third, Viewpoint Persistence: camera azimuth, elevation, and zoom coordinates are strictly preserved across filter updates using Plotly's uirevision property."

[VISUAL ACTION: Highlight Geographic Coordinate Retention, Aspect Ratio Correction (1 km metric proportionality), and Viewpoint Persistence on the cards]""",

    # SLIDE 7: Multi-Layer Geophysics
    """[SLIDE 7 | TIME: 4:00 - 4:45 | TARGET: 45 SECONDS]

SPEECH (EN):
"A central contribution of Seismic Viewer is fusing multiple geophysical layers in the exact same 3D coordinate space.

First, 3D Vp/Vs Seismic Tomography: we convert tomographic velocity ratios into a 3D volumetric grid using Inverse Distance Weighting with 16 nearest neighbours, applying a Gaussian spatial decay that smoothly blends toward the regional crustal baseline outside sampled points.
Second, GNSS Ground Uplift Columns: we ingest GPS deformation time-series across Campi Flegrei, rendering explicit 3D vertical displacement columns directly at station locations.
By seeing these layers together, researchers can directly correlate surface ground swelling with underlying fluid pockets and earthquake swarms."

[VISUAL ACTION: Point to the 3D Vp/Vs tomography volume on Card 1 and the GNSS vertical uplift columns on Card 2]""",

    # SLIDE 8: 4D Temporal Animation & Sequence Slicing
    """[SLIDE 8 | TIME: 4:45 - 5:25 | TARGET: 40 SECONDS]

SPEECH (EN):
"Earthquake swarms migrate dynamically over time. Our 4D engine provides sequence slicing across three distinct analytical playback modes:

Accumulation Mode, showing cumulative seismic buildup from the first event to track stress release over time;
Sliding Window Mode, which moves a window of width W across time—typically one frame per day—to isolate active hypocentral migration fronts;
and Block Mode, moving a window with a custom step to allow flexible temporal downsampling and multi-day epoch comparisons.
At the bottom, an Interactive Timeline Histogram couples chronological navigation with real-time tracking of active event counts, peak magnitude, and cumulative seismic energy release."

[VISUAL ACTION: Walk through the three modes (Accumulation, Sliding Window, Block Mode) and point to the interactive timeline histogram and energy meter at the bottom]""",

    # SLIDE 9: Algorithmic Innovation — Cylinders & Label Stability
    """[SLIDE 9 | TIME: 5:25 - 6:05 | TARGET: 40 SECONDS]

SPEECH (EN):
"Next, we solved two practical algorithmic challenges:

First, Why Cylinder Partitioning? Caldera ring faults curve along the volcanic perimeter, so attempting to fit a single global plane produces large geometric errors. Working directly with Giordano and Alfonsi, we replicated their spatial segmentation methodology, partitioning the crust into vertical cylindrical columns to isolate locally planar fault segments and validate our automated fitting against their published results.
Second, Global Label Matching: in 4D animations, standard clustering randomly reassigns cluster IDs between time steps, causing chaotic color flickering. We developed a historical majority voting algorithm that maintains stable cluster identities across frames.
Finally, an asynchronous worker pool of 8 threads parallelizes cylinder fitting across animation steps."

[VISUAL ACTION: Point to Why Cylinder Partitioning on Card 1, Global Label Matching on Card 2, and the 8-worker thread pool on Card 3]""",

    # SLIDE 10: Machine Learning Clustering Engine
    """[SLIDE 10 | TIME: 6:05 - 6:50 | TARGET: 45 SECONDS]

SPEECH (EN):
"To automatically isolate fault structures from thousands of earthquakes, we designed a flexible clustering engine.

We initially evaluated K-Means, but it fails in seismology: it forces every point into a cluster and assumes spherical shapes. In reality, geological faults are elongated, and tectonic background noise surrounds them.
Therefore, we adopted DBSCAN: it detects arbitrary, non-planar fault shapes without requiring a predefined cluster count, and correctly isolates background noise as label minus one. We automated radius tuning by estimating epsilon from the k-nearest-neighbour knee using the Kneedle algorithm.
We also retained MiniBatchKMeans for rapid centroid preview, and introduced a Hierarchical Cluster Lock mechanism: researchers can selectively lock confirmed faults and re-cluster only unresolved swarms with refined local parameters."

[VISUAL ACTION: Explain why K-Means failed and DBSCAN succeeded, then highlight the Hierarchical Sub-Clustering and Cluster Lock mechanism on the main card]""",

    # SLIDE 11: Structural Geology & Fault Plane Fitting
    """[SLIDE 11 | TIME: 6:50 - 7:35 | TARGET: 45 SECONDS]

SPEECH (EN):
"From clustered hypocentres, the system mathematically reconstructs 3D fault planes in two steps:

Why not just fit a plane using SVD directly? Because outlier points or secondary micro-cracks will pull the plane off its true orientation.
So, Step 1 is RANSAC Inlier Consensus: it tests random 3-point combinations over 100 iterations to find the planar consensus and discard non-planar noise.
Step 2 applies Singular Value Decomposition strictly to those centered inliers. The fault normal vector comes directly from the eigenvector of the smallest singular value.
From this normal vector, the system calculates geological Strike, Dip, Dip Direction, and Planarity, rendering interactive 3D rectangular plane meshes with hover vectors."

[VISUAL ACTION: Walk through Step 1 (RANSAC Inlier Consensus), Step 2 (SVD & Geological Angles), and point to the Strike, Dip, and Planarity outputs]""",

    # SLIDE 12: Data Engineering Pipeline
    """[SLIDE 12 | TIME: 7:35 - 8:20 | TARGET: 45 SECONDS]

SPEECH (EN):
"Looking at the data engineering pipeline:

Stage 1 handles Heterogeneous Ingestion: querying the INGV GOSSIP REST API in yearly chunks alongside local offline CSV parsing, supported by automatic alias detection for column headers.
Stage 2 enforces Quality Sanitisation and our Missing Depth Principle: earthquakes without a recorded depth are retained in 2D views for catalogue completeness, but strictly excluded from 3D views and clustering to prevent artificial planes at sea level. Physical bounds checks restrict depths between minus 10 and 800 kilometres.
Finally, in the bottom hero card, we replaced heavy JSON with our High-Performance Float32 Binary Protocol.
Streaming raw contiguous binary buffers slashed payload size by nearly 80%—from 4.75 megabytes to 976 kilobytes—and dropped browser decode time to less than 0.01 milliseconds."

[VISUAL ACTION: Trace through the 4 top cards (Ingestion, Missing Depth Principle, Physical Bounds, WGS84) and highlight the -80% size and <0.01ms badges on the bottom card]""",

    # SLIDE 13: Results & Quantitative Validation Benchmarks
    """[SLIDE 13 | TIME: 8:20 - 9:05 | TARGET: 45 SECONDS]

SPEECH (EN):
"We evaluated the system thoroughly across four benchmark dimensions:

First, Structural Accuracy in Monte Carlo simulations across 500 synthetic faults with noise up to 100 metres and 30% outlier contamination: our RANSAC-SVD pipeline achieved a Mean Absolute Error under 0.51 degrees for Dip, and under 0.93 degrees for Strike.
Second, Real Catalogue Testing on 3,225 INGV events from Campi Flegrei executed clustering and plane extraction in under 42 milliseconds.
Third, Algorithmic Throughput: cylinder grid partitioning runs in 36.5 milliseconds, and stress tests with 50,000 synthetic hypocentres maintained 55 to 60 frames per second.
And fourth, Serialization & Memory: browser decode dropped from 39.7 to under 0.01 milliseconds, and FIFO session management holds RAM strictly below 500 megabytes."

[VISUAL ACTION: Walk through the 4 benchmark cards: Monte Carlo MAE badges (<0.51° Dip, <0.93° Strike), Real Catalogue 42ms speed, 60 FPS stress test, and <500 MB RAM]""",

    # SLIDE 14: Conclusions & Future Work
    """[SLIDE 14 | TIME: 9:05 - 9:35 | TARGET: 30 SECONDS]

SPEECH (EN):
"To conclude, Seismic Viewer delivers on three validated contributions:
A High-Performance Architecture pairing a pywebview client with an asynchronous FastAPI server and Float32 streaming;
Multi-Layer Visualisation integrating 4D animations with 3D tomography and GNSS uplift;
and Robust Structural Analysis benchmarked with sub-degree angular accuracy.

Looking forward, future extensions will focus on Real-Time Surveillance Feeds with WebSockets and SeedLink, WebGPU compute shaders for client-side DBSCAN re-clustering, and 3D focal mechanism 'beachballs' for kinematic seismotectonics."

[VISUAL ACTION: Highlight the 3 Validated Contributions on the left, the 3 Future Extensions on the right, and then transition to Slide 15]""",

    # SLIDE 15: Closing & Thank You
    """[SLIDE 15 | TIME: 9:35 - 9:45 | TARGET: 10 SECONDS]

SPEECH (EN):
"Thank you very much for your time and attention.
I would be glad to answer any questions from the Committee."

[VISUAL ACTION: Smile warmly, acknowledge Professor Avola and Dr. Bailo, thank the entire committee, and open the floor for questions]"""
]

pptx_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Seismic Viewer - Thesis Defense.pptx')
prs = pptx.Presentation(pptx_path)
assert len(prs.slides) == len(SPEECH_NOTES_SIMPLE_ENGLISH), f"Mismatch: {len(prs.slides)} slides vs {len(SPEECH_NOTES_SIMPLE_ENGLISH)} notes"

for i, slide in enumerate(prs.slides):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = SPEECH_NOTES_SIMPLE_ENGLISH[i]

prs.save(pptx_path)
print("Successfully injected all 15 simplified notes based directly on slide content into 'Seismic Viewer - Thesis Defense.pptx'!")
