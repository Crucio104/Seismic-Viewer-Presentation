import pptx
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt, Emu
import os
import win32com.client

script_dir = os.path.dirname(os.path.abspath(__file__))
pptx_path = os.path.join(script_dir, "Seismic Viewer - Thesis Defense.pptx")

# Palette
BORDEAUX = RGBColor(130, 36, 51)       # #822433
BORDEAUX_DARK = RGBColor(100, 20, 35)  # #641423
BORDEAUX_TINT = RGBColor(246, 237, 239)# #F6EDEF
BORDEAUX_LINE = RGBColor(180, 60, 78)  # #B43C4E
BEIGE_BG = RGBColor(250, 248, 242)
CARD_BG = RGBColor(255, 255, 255)
CARD_LINE = RGBColor(228, 222, 212)
TEXT_HEADING = RGBColor(28, 24, 25)
TEXT_BODY = RGBColor(48, 42, 44)
TEXT_MUTED = RGBColor(115, 105, 103)
CARD_HEADER_BG = RGBColor(248, 244, 238)

prs = pptx.Presentation(pptx_path)
slide4 = prs.slides[3]

# Let's inspect shapes to keep: Slide titles (Shape 8, 9), Sapienza logo (Shape 22), SapienzaTopStripe (Shape 23)
# Background cards (Shape 0, 1)
# We can clear old text boxes on Slide 4 and rebuild Card 1 (Left: Requirements) and Card 2 (Right: Timeline)

# Update subtitle (Shape 9)
slide4.shapes[9].text_frame.text = "Requirements & Iterative Internship Development Timeline"
slide4.shapes[9].text_frame.paragraphs[0].font.name = "Arial"
slide4.shapes[9].text_frame.paragraphs[0].font.size = Pt(20)
slide4.shapes[9].text_frame.paragraphs[0].font.bold = True
slide4.shapes[9].text_frame.paragraphs[0].font.color.rgb = BORDEAUX

# Delete old intermediate text shapes (Shapes 2 to 7, 10 to 21)
# Let's collect shape IDs to remove
to_remove = []
for s in slide4.shapes:
    if s.shape_id in [149, 150, 151, 152, 153, 154, 157, 158, 159, 160, 161, 162, 163, 164, 165, 166, 167, 168]:
        to_remove.append(s)

for s in to_remove:
    sp = s._element
    sp.getparent().remove(sp)

print(f"Removed {len(to_remove)} legacy shapes from Slide 4.")

# Ensure Card 0 and Card 1 have proper clean card styling
c0 = slide4.shapes[0]
c1 = slide4.shapes[1]
c0.left = 609600; c0.top = 1350000; c0.width = 5372100; c0.height = 5150000
c1.left = 6210300; c1.top = 1350000; c1.width = 5372100; c1.height = 5150000

# Set fill and border
for card in [c0, c1]:
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_LINE
    card.line.width = Pt(1.5)

# --- LEFT CARD: System Requirements & Constraints ---
# Header Box
h0 = slide4.shapes.add_textbox(c0.left + Inches(0.2), c0.top + Inches(0.18), c0.width - Inches(0.4), Inches(0.45))
tf0 = h0.text_frame
tf0.word_wrap = True
tf0.margin_left = tf0.margin_right = tf0.margin_top = tf0.margin_bottom = 0
p = tf0.paragraphs[0]
p.text = "System Requirements & Constraints"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = BORDEAUX

# Body Box for Requirements
b0 = slide4.shapes.add_textbox(c0.left + Inches(0.2), c0.top + Inches(0.65), c0.width - Inches(0.4), Inches(4.3))
tf_b0 = b0.text_frame
tf_b0.word_wrap = True
tf_b0.margin_left = tf_b0.margin_right = tf_b0.margin_top = tf_b0.margin_bottom = 0

req_items = [
    ("Functional Requirements (FR)", True, BORDEAUX, 11),
    ("• FR1 — Multi-Modal 4D Navigation: Seamless toggle between 2D maps, 3D point cloud, and dynamic 4D time playback.", False, TEXT_BODY, 9.5),
    ("• FR2 — Automated Fault Extraction: Mathematical plane fitting (RANSAC + SVD) with Strike & Dip angles and planarity.", False, TEXT_BODY, 9.5),
    ("• FR3 — Multi-Layer Geophysics: Volumetric Vp/Vs tomography, GNSS ground uplift, and seismic station positions.", False, TEXT_BODY, 9.5),
    ("• FR4 — Heterogeneous Ingestion: Online queries to INGV GOSSIP API and offline local CSV parsing.", False, TEXT_BODY, 9.5),
    ("Non-Functional Requirements (NFR)", True, BORDEAUX, 11),
    ("• NFR1 — Real-Time Rendering: 50–60 FPS with up to 50,000 seismic hypocentres via WebGL GPU instancing.", False, TEXT_BODY, 9.5),
    ("• NFR2 — Sub-Second Latency: Interactive clustering recalculation under 500 ms for slider response.", False, TEXT_BODY, 9.5),
    ("• NFR3 — Bounded Footprint: Native desktop pywebview; memory footprint strictly under 500 MB.", False, TEXT_BODY, 9.5),
    ("• NFR4 — Robust Sanitisation: Safe exclusion of missing depth events from 3D to prevent planar artifacts.", False, TEXT_BODY, 9.5),
]

for idx, (text, is_bold, color, size) in enumerate(req_items):
    p = tf_b0.paragraphs[0] if idx == 0 else tf_b0.add_paragraph()
    p.text = text
    p.font.name = "Arial"
    p.font.size = Pt(size)
    p.font.bold = is_bold
    p.font.color.rgb = color
    p.space_after = Pt(4 if not is_bold else 2)
    p.space_before = Pt(6 if is_bold and idx > 0 else 0)


# --- RIGHT CARD: Iterative Implementation Timeline ---
# Header Box
h1 = slide4.shapes.add_textbox(c1.left + Inches(0.2), c1.top + Inches(0.18), c1.width - Inches(0.4), Inches(0.45))
tf1 = h1.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_right = tf1.margin_top = tf1.margin_bottom = 0
p = tf1.paragraphs[0]
p.text = "Implementation Timeline (Internship Phases)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = BORDEAUX

# Subtitle Box
sub1 = slide4.shapes.add_textbox(c1.left + Inches(0.2), c1.top + Inches(0.55), c1.width - Inches(0.4), Inches(0.3))
tf_sub1 = sub1.text_frame
tf_sub1.word_wrap = True
tf_sub1.margin_left = tf_sub1.margin_right = tf_sub1.margin_top = tf_sub1.margin_bottom = 0
p = tf_sub1.paragraphs[0]
p.text = "Iterative development shaped by continuous feedback from INGV & Roma Tre"
p.font.name = "Arial"
p.font.size = Pt(9.5)
p.font.italic = True
p.font.color.rgb = TEXT_MUTED

# Timeline items
b1 = slide4.shapes.add_textbox(c1.left + Inches(0.2), c1.top + Inches(0.9), c1.width - Inches(0.4), Inches(4.1))
tf_b1 = b1.text_frame
tf_b1.word_wrap = True
tf_b1.margin_left = tf_b1.margin_right = tf_b1.margin_top = tf_b1.margin_bottom = 0

timeline_items = [
    ("Phase 1: 3D Prototype & CSV Ingestion", "Initial 3D hypocentral point cloud visualization and local CSV file parsing engine."),
    ("Phase 2: K-Means Clustering & Native Bridge", "First spatial clustering implementation and two-way bridge between Python and WebView."),
    ("Phase 3: UI Redesign & INGV GOSSIP API", "Modern analytical interface and automated online catalogue query directly from INGV."),
    ("Phase 4: DBSCAN Knee Tuning & Cylinders", "Automatic knee detection for eps, 4D cylinder partitioning, and confirmed cluster locking."),
    ("Phase 5: Multi-Layer Geophysical Fusion", "Volumetric 3D tomography (Vp/Vs), GNSS deformation surface, and station markers."),
    ("Phase 6: Binary Streaming & Async Architecture", "Float32 contiguous binary streaming (-80% payload) and 8-worker ThreadPool backend."),
]

for idx, (title, desc) in enumerate(timeline_items):
    p = tf_b1.paragraphs[0] if idx == 0 else tf_b1.add_paragraph()
    run_badge = p.add_run()
    run_badge.text = f"P{idx+1} "
    run_badge.font.name = "Arial"
    run_badge.font.size = Pt(9.5)
    run_badge.font.bold = True
    run_badge.font.color.rgb = BORDEAUX

    run_title = p.add_run()
    run_title.text = title
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10)
    run_title.font.bold = True
    run_title.font.color.rgb = TEXT_HEADING

    p2 = tf_b1.add_paragraph()
    run_desc = p2.add_run()
    run_desc.text = desc
    run_desc.font.name = "Arial"
    run_desc.font.size = Pt(9.2)
    run_desc.font.color.rgb = TEXT_BODY
    p2.space_after = Pt(5)

prs.save(pptx_path)
print("Updated Slide 4 successfully with System Requirements & Implementation Timeline!")

# Export slide 4 to check visually
try:
    ppt = win32com.client.Dispatch("PowerPoint.Application")
    pres = ppt.Presentations.Open(pptx_path, WithWindow=False)
    out_path = os.path.join(script_dir, "exported_slides", "slide_04.png")
    pres.Slides[4].Export(out_path, "PNG", 1920, 1080)
    pres.Close()
    ppt.Quit()
    print(f"Exported updated slide 4 to {out_path}")
except Exception as e:
    print(f"PowerPoint export failed: {e}")
