import pptx
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os
import shutil
import win32com.client

script_dir = os.path.dirname(os.path.abspath(__file__))
pptx_path = os.path.join(script_dir, "Seismic Viewer - Thesis Defense.pptx")
desktop_pptx_path = r"C:\Users\aless\Desktop\Seismic Viewer - Thesis Defense.pptx"

# Palette
BORDEAUX = RGBColor(130, 36, 51)       # #822433
BORDEAUX_DARK = RGBColor(100, 20, 35)  # #641423
BORDEAUX_TINT = RGBColor(246, 237, 239)# #F6EDEF
BORDEAUX_LINE = RGBColor(180, 60, 78)  # #B43C4E
BEIGE_BG = RGBColor(250, 248, 242)
CARD_BG = RGBColor(255, 255, 255)
CARD_LINE = RGBColor(228, 222, 212)
CARD_HEADER_BG = RGBColor(248, 244, 238)
TEXT_HEADING = RGBColor(28, 24, 25)
TEXT_BODY = RGBColor(48, 42, 44)
TEXT_MUTED = RGBColor(115, 105, 103)
PHASE_BG = RGBColor(253, 251, 248)
TAG_BG = RGBColor(240, 235, 230)

prs = pptx.Presentation(pptx_path)
slide4 = prs.slides[3]

# 1. Update Subtitle
slide4.shapes[3].text_frame.text = "System Requirements & Iterative Internship Roadmap"
slide4.shapes[3].text_frame.paragraphs[0].font.name = "Arial"
slide4.shapes[3].text_frame.paragraphs[0].font.size = Pt(20)
slide4.shapes[3].text_frame.paragraphs[0].font.bold = True
slide4.shapes[3].text_frame.paragraphs[0].font.color.rgb = BORDEAUX

# 2. Cleanup all previous custom shapes on Slide 4
shapes_to_keep_ids = [slide4.shapes[2].shape_id, slide4.shapes[3].shape_id, slide4.shapes[4].shape_id, slide4.shapes[5].shape_id]
to_remove = [s for s in slide4.shapes if s.shape_id not in shapes_to_keep_ids]

for s in to_remove:
    sp = s._element
    sp.getparent().remove(sp)

print(f"Cleaned up {len(to_remove)} legacy shapes from Slide 4.")

# Layout Dimensions
TOTAL_W = 10972800  # 12 inches
LEFT_MARGIN = 609600

# =========================================================================
# TOP ROW: System Requirements (FR on Left, NFR on Right)
# =========================================================================
TOP_Y = 1320000
TOP_H = 2600000
GAP_X = 232800
CARD_W = (TOTAL_W - GAP_X) // 2  # 5370000

# ----------------- TOP-LEFT CARD: Functional Requirements -----------------
fr_card = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, TOP_Y, CARD_W, TOP_H)
fr_card.fill.solid()
fr_card.fill.fore_color.rgb = CARD_BG
fr_card.line.color.rgb = CARD_LINE
fr_card.line.width = Pt(1.5)

# Header banner FR
fr_header = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, TOP_Y, CARD_W, Inches(0.48))
fr_header.fill.solid()
fr_header.fill.fore_color.rgb = BORDEAUX_TINT
fr_header.line.fill.background()

tf_fr_h = fr_header.text_frame
tf_fr_h.word_wrap = True
tf_fr_h.margin_left = Inches(0.2)
tf_fr_h.margin_top = Inches(0.06)
p = tf_fr_h.paragraphs[0]
r1 = p.add_run()
r1.text = "Functional Requirements (FR)"
r1.font.name = "Arial"
r1.font.size = Pt(12)
r1.font.bold = True
r1.font.color.rgb = BORDEAUX

r2 = p.add_run()
r2.text = "   |   Core scientific capabilities & user workflows"
r2.font.name = "Arial"
r2.font.size = Pt(8.5)
r2.font.italic = True
r2.font.color.rgb = TEXT_MUTED

# Body FR
fr_body = slide4.shapes.add_textbox(LEFT_MARGIN + Inches(0.2), TOP_Y + Inches(0.55), CARD_W - Inches(0.4), TOP_H - Inches(0.6))
tf_fr_b = fr_body.text_frame
tf_fr_b.word_wrap = True
tf_fr_b.margin_left = tf_fr_b.margin_right = tf_fr_b.margin_top = tf_fr_b.margin_bottom = 0

fr_bullets = [
    ("FR1 — Multi-Modal 4D Navigation:", " Real-time toggle across synchronized 2D map, 3D WebGL point cloud, and 4D temporal playback."),
    ("FR2 — Automated Fault Plane Fitting:", " Mathematical extraction (RANSAC + SVD) computing Strike, Dip, and planarity ratio."),
    ("FR3 — Multi-Layer Geophysical Fusion:", " Volumetric 3D tomography (Vp/Vs), GNSS surface uplift deformation, and seismic station positions."),
    ("FR4 — Dual Catalogue Ingestion:", " Automated queries to the INGV GOSSIP REST API and local offline CSV parsing engine.")
]

for idx, (tag, desc) in enumerate(fr_bullets):
    p = tf_fr_b.paragraphs[0] if idx == 0 else tf_fr_b.add_paragraph()
    r_tag = p.add_run()
    r_tag.text = f"• {tag}"
    r_tag.font.name = "Arial"
    r_tag.font.size = Pt(9.5)
    r_tag.font.bold = True
    r_tag.font.color.rgb = BORDEAUX_DARK
    
    r_desc = p.add_run()
    r_desc.text = desc
    r_desc.font.name = "Arial"
    r_desc.font.size = Pt(9.0)
    r_desc.font.color.rgb = TEXT_BODY
    p.space_after = Pt(6)


# ----------------- TOP-RIGHT CARD: Non-Functional Requirements -----------------
RIGHT_X = LEFT_MARGIN + CARD_W + GAP_X
nfr_card = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, RIGHT_X, TOP_Y, CARD_W, TOP_H)
nfr_card.fill.solid()
nfr_card.fill.fore_color.rgb = CARD_BG
nfr_card.line.color.rgb = CARD_LINE
nfr_card.line.width = Pt(1.5)

# Header banner NFR
nfr_header = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, RIGHT_X, TOP_Y, CARD_W, Inches(0.48))
nfr_header.fill.solid()
nfr_header.fill.fore_color.rgb = CARD_HEADER_BG
nfr_header.line.fill.background()

tf_nfr_h = nfr_header.text_frame
tf_nfr_h.word_wrap = True
tf_nfr_h.margin_left = Inches(0.2)
tf_nfr_h.margin_top = Inches(0.06)
p = tf_nfr_h.paragraphs[0]
r1 = p.add_run()
r1.text = "Non-Functional Requirements (NFR)"
r1.font.name = "Arial"
r1.font.size = Pt(12)
r1.font.bold = True
r1.font.color.rgb = BORDEAUX

r2 = p.add_run()
r2.text = "   |   Performance limits, footprint & scientific integrity"
r2.font.name = "Arial"
r2.font.size = Pt(8.5)
r2.font.italic = True
r2.font.color.rgb = TEXT_MUTED

# Body NFR
nfr_body = slide4.shapes.add_textbox(RIGHT_X + Inches(0.2), TOP_Y + Inches(0.55), CARD_W - Inches(0.4), TOP_H - Inches(0.6))
tf_nfr_b = nfr_body.text_frame
tf_nfr_b.word_wrap = True
tf_nfr_b.margin_left = tf_nfr_b.margin_right = tf_nfr_b.margin_top = tf_nfr_b.margin_bottom = 0

nfr_bullets = [
    ("NFR1 — Real-Time Rendering:", " Sustained 50–60 FPS with up to 50,000 seismic hypocentres via WebGL GPU instancing."),
    ("NFR2 — Sub-Second Latency:", " Interactive clustering recomputation under 500 ms for instantaneous parameter tuning."),
    ("NFR3 — Bounded Footprint:", " Native cross-platform desktop wrapper (pywebview) strictly under 500 MB RAM consumption."),
    ("NFR4 — Robust Sanitisation:", " Automatic exclusion of zero-depth / corrupted events to avoid planar fitting distortions.")
]

for idx, (tag, desc) in enumerate(nfr_bullets):
    p = tf_nfr_b.paragraphs[0] if idx == 0 else tf_nfr_b.add_paragraph()
    r_tag = p.add_run()
    r_tag.text = f"• {tag}"
    r_tag.font.name = "Arial"
    r_tag.font.size = Pt(9.5)
    r_tag.font.bold = True
    r_tag.font.color.rgb = BORDEAUX_DARK
    
    r_desc = p.add_run()
    r_desc.text = desc
    r_desc.font.name = "Arial"
    r_desc.font.size = Pt(9.0)
    r_desc.font.color.rgb = TEXT_BODY
    p.space_after = Pt(6)


# =========================================================================
# BOTTOM ROW: Internship Development Timeline (6 Horizontal Phases)
# =========================================================================
BOTTOM_Y = 4070000
BOTTOM_H = 2500000

# Container card
timeline_card = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN, BOTTOM_Y, TOTAL_W, BOTTOM_H)
timeline_card.fill.solid()
timeline_card.fill.fore_color.rgb = CARD_BG
timeline_card.line.color.rgb = CARD_LINE
timeline_card.line.width = Pt(1.5)

# Header inside timeline card
t_header = slide4.shapes.add_textbox(LEFT_MARGIN + Inches(0.2), BOTTOM_Y + Inches(0.12), TOTAL_W - Inches(0.4), Inches(0.35))
tf_th = t_header.text_frame
tf_th.word_wrap = True
tf_th.margin_left = tf_th.margin_right = tf_th.margin_top = tf_th.margin_bottom = 0
p = tf_th.paragraphs[0]
r_th = p.add_run()
r_th.text = "Internship Implementation Timeline"
r_th.font.name = "Arial"
r_th.font.size = Pt(12.5)
r_th.font.bold = True
r_th.font.color.rgb = BORDEAUX

r_th_sub = p.add_run()
r_th_sub.text = "   —   6 iterative phases guided by continuous geophysicist feedback (INGV & Roma Tre)"
r_th_sub.font.name = "Arial"
r_th_sub.font.size = Pt(9.2)
r_th_sub.font.italic = True
r_th_sub.font.color.rgb = TEXT_MUTED

# Connecting timeline horizontal bar
bar_y = BOTTOM_Y + Inches(0.58)
bar_line = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN + Inches(0.5), bar_y, TOTAL_W - Inches(1.0), Pt(2.5))
bar_line.fill.solid()
bar_line.fill.fore_color.rgb = RGBColor(215, 205, 195)
bar_line.line.fill.background()

# 6 Phase Columns
phases = [
    ("Phase 1", "3D Prototype & CSV", "Hypocentre 3D cloud & local CSV parser.", "Three.js & WebGL"),
    ("Phase 2", "K-Means & Bridge", "Spatial clustering & Python-WebView bridge.", "pywebview & IPC"),
    ("Phase 3", "UI Redesign & API", "Analytical dashboard & INGV GOSSIP API.", "Leaflet & REST"),
    ("Phase 4", "DBSCAN & Cylinders", "Knee tuning, 4D cylinder slicing & locking.", "k-NN Knee & 4D"),
    ("Phase 5", "Geophysical Layers", "Vp/Vs tomography, GNSS surface & stations.", "Mesh & Tomography"),
    ("Phase 6", "Streaming & Async", "Float32 streaming (-80% payload) & worker pool.", "Binary & ThreadPool")
]

num_phases = len(phases)
col_gap = Inches(0.12)
total_avail_w = TOTAL_W - Inches(0.4)
phase_w = int((total_avail_w - (num_phases - 1) * col_gap) / num_phases)
col_top = BOTTOM_Y + Inches(0.72)
col_h = BOTTOM_H - Inches(0.84)

for i, (p_tag, p_title, p_desc, p_stack) in enumerate(phases):
    col_left = int(LEFT_MARGIN + Inches(0.2) + i * (phase_w + col_gap))
    
    # Milestone connector dot on the timeline bar
    dot_size = Inches(0.18)
    dot_x = int(col_left + phase_w // 2 - dot_size // 2)
    dot_y = int(bar_y - Inches(0.075))
    dot = slide4.shapes.add_shape(MSO_SHAPE.OVAL, dot_x, dot_y, dot_size, dot_size)
    dot.fill.solid()
    dot.fill.fore_color.rgb = BORDEAUX if i < 5 else BORDEAUX_DARK
    dot.line.color.rgb = CARD_BG
    dot.line.width = Pt(2)
    
    # Phase mini-card
    p_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col_left, col_top, phase_w, col_h)
    p_box.fill.solid()
    p_box.fill.fore_color.rgb = PHASE_BG if i < 5 else RGBColor(254, 248, 249)
    p_box.line.color.rgb = BORDEAUX_LINE if i == 5 else CARD_LINE
    p_box.line.width = Pt(1.5 if i == 5 else 1)
    
    # Phase badge inside box
    badge_h = Inches(0.24)
    badge = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, col_left, col_top, phase_w, badge_h)
    badge.fill.solid()
    badge.fill.fore_color.rgb = BORDEAUX if i == 5 else BORDEAUX_TINT
    badge.line.fill.background()
    tf_b = badge.text_frame
    tf_b.word_wrap = True
    tf_b.margin_top = Inches(0.02)
    tf_b.margin_bottom = tf_b.margin_left = tf_b.margin_right = 0
    p = tf_b.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = p_tag.upper()
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255) if i == 5 else BORDEAUX
    
    # Phase title, desc, and stack
    content_box = slide4.shapes.add_textbox(col_left + Inches(0.08), col_top + badge_h + Inches(0.06), phase_w - Inches(0.16), col_h - badge_h - Inches(0.1))
    tf_c = content_box.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
    
    # Title
    p_t = tf_c.paragraphs[0]
    p_t.text = p_title
    p_t.font.name = "Arial"
    p_t.font.size = Pt(9.0)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_HEADING
    p_t.space_after = Pt(3)
    
    # Desc
    p_d = tf_c.add_paragraph()
    p_d.text = p_desc
    p_d.font.name = "Arial"
    p_d.font.size = Pt(8.2)
    p_d.font.color.rgb = TEXT_BODY
    p_d.space_after = Pt(5)
    
    # Tech / Deliverable pill tag
    p_tech = tf_c.add_paragraph()
    r_tech_lbl = p_tech.add_run()
    r_tech_lbl.text = "Focus: "
    r_tech_lbl.font.name = "Arial"
    r_tech_lbl.font.size = Pt(7.8)
    r_tech_lbl.font.bold = True
    r_tech_lbl.font.color.rgb = BORDEAUX
    
    r_tech_val = p_tech.add_run()
    r_tech_val.text = p_stack
    r_tech_val.font.name = "Arial"
    r_tech_val.font.size = Pt(7.8)
    r_tech_val.font.color.rgb = TEXT_MUTED

# Save to workspace pptx
prs.save(pptx_path)
print("Saved refined Slide 4 to workspace pptx successfully!")

# Also copy to desktop pptx so both files are strictly synchronized!
try:
    shutil.copy2(pptx_path, desktop_pptx_path)
    print(f"Synchronized copy to {desktop_pptx_path} successfully!")
except Exception as e:
    print(f"Notice: Could not copy to desktop: {e}")

# Export slide 4 to PNG for visual inspection
try:
    ppt = win32com.client.Dispatch("PowerPoint.Application")
    pres = ppt.Presentations.Open(pptx_path, WithWindow=False)
    out_dir = os.path.join(script_dir, "exported_slides")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "slide_04.png")
    pres.Slides.Item(4).Export(out_path, "PNG", 1920, 1080)
    pres.Close()
    ppt.Quit()
    print(f"Exported refined slide 4 to {out_path}")
except Exception as e:
    print(f"PowerPoint export note: {e}")
