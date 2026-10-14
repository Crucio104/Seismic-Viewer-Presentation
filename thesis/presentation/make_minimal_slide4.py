import pptx
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os
import shutil
import win32com.client

script_dir = os.path.dirname(os.path.abspath(__file__))
pptx_path = os.path.join(script_dir, "Seismic Viewer - Thesis Defense.pptx")
desktop_pptx_path = r"C:\Users\aless\Desktop\Seismic Viewer - Thesis Defense.pptx"

# Exact Palette from Slides 2, 5, 7, 12
BORDEAUX = RGBColor(130, 36, 51)          # #822433 (Sapienza primary)
BORDEAUX_DARK = RGBColor(100, 20, 35)     # #641423
BORDEAUX_TINT_BG = RGBColor(251, 244, 243)# #FBF4F3 (soft bordeaux card fill from Slide 7)
BORDEAUX_TINT_LN = RGBColor(232, 213, 216)# #E8D5D8 (soft bordeaux border from Slide 7)

BLUE_TINT_BG = RGBColor(242, 246, 250)    # #F2F6FA (soft blue card fill from Slide 7 & 12)
BLUE_TINT_LN = RGBColor(208, 220, 229)    # #D0DCE5 (soft blue border from Slide 7 & 12)

BEIGE_TINT_BG = RGBColor(253, 248, 242)   # #FDF8F2 (warm ivory card fill from Slide 7 & 12)
BEIGE_TINT_LN = RGBColor(235, 215, 197)   # #EBD7C5 (warm ivory border from Slide 7 & 12)

NAVY_CATEGORY = RGBColor(32, 64, 96)      # #204060
TEXT_HEADING = RGBColor(28, 24, 25)       # #1C1819
TEXT_BODY = RGBColor(48, 42, 44)          # #302A2C
TEXT_MUTED = RGBColor(115, 105, 103)      # #736967
TIMELINE_LINE_CLR = RGBColor(218, 210, 200)

prs = pptx.Presentation(pptx_path)
slide4 = prs.slides[3]

# Identify Sapienza logo shape
logo_shape = None
for s in slide4.shapes:
    if s.name.startswith("Picture") or s.shape_id == 169:
        logo_shape = s
        break

# Clear all shapes on Slide 4 except Logo if found
to_remove = []
for s in slide4.shapes:
    if s != logo_shape:
        to_remove.append(s)

for s in to_remove:
    sp = s._element
    sp.getparent().remove(sp)

print(f"Removed {len(to_remove)} shapes to rebuild minimalist slide 4.")

# 1. Top Bordeaux Stripe (identical to other slides)
top_stripe = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 12192000, 76200)
top_stripe.fill.solid()
top_stripe.fill.fore_color.rgb = BORDEAUX
top_stripe.line.fill.background()

# 2. Sapienza Logo
if logo_shape is None:
    logo_file = os.path.join(script_dir, "sapienza_logo.png")
    if os.path.exists(logo_file):
        slide4.shapes.add_picture(logo_file, 9750000, 380000, width=1900000, height=570000)
else:
    logo_shape.left = 9750000
    logo_shape.top = 380000
    logo_shape.width = 1900000
    logo_shape.height = 570000

# 3. Category Header
cat_box = slide4.shapes.add_textbox(609600, 560000, 9000000, 180000)
tf_cat = cat_box.text_frame
tf_cat.word_wrap = True
tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
p = tf_cat.paragraphs[0]
p.text = "SYSTEM DESIGN & SPECIFICATIONS"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = NAVY_CATEGORY

# 4. Slide Title (Single line, 22pt bold)
title_box = slide4.shapes.add_textbox(609600, 780000, 9000000, 420000)
tf_title = title_box.text_frame
tf_title.word_wrap = False
tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
p = tf_title.paragraphs[0]
p.text = "System Requirements & Iterative Internship Roadmap"
p.font.name = "Arial"
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = BORDEAUX

# Layout Metrics
TOTAL_W = 10972800  # 12 inches
LEFT_MARGIN = 609600

# =========================================================================
# TOP ROW: System Requirements Cards (Same style as Slides 2, 7 & 12)
# =========================================================================
TOP_Y = 1350000
TOP_H = 2600000
GAP_X = Inches(0.24)
CARD_W = int((TOTAL_W - GAP_X) / 2)

# ----------------- TOP-LEFT CARD: Functional Requirements (FR) -----------------
# Seamless Rounded Rectangle with soft bordeaux tint (matching Slide 7 Card 0)
fr_card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, LEFT_MARGIN, TOP_Y, CARD_W, TOP_H)
fr_card.fill.solid()
fr_card.fill.fore_color.rgb = BORDEAUX_TINT_BG
fr_card.line.color.rgb = BORDEAUX_TINT_LN
fr_card.line.width = Pt(1.2)

# Text inside FR Card
fr_text = slide4.shapes.add_textbox(LEFT_MARGIN + Inches(0.22), TOP_Y + Inches(0.16), CARD_W - Inches(0.44), TOP_H - Inches(0.32))
tf_fr = fr_text.text_frame
tf_fr.word_wrap = True
tf_fr.margin_left = tf_fr.margin_right = tf_fr.margin_top = tf_fr.margin_bottom = 0

# Eyebrow tag
p_eye = tf_fr.paragraphs[0]
p_eye.text = "CORE SCIENTIFIC WORKFLOWS"
p_eye.font.name = "Arial"
p_eye.font.size = Pt(9.2)
p_eye.font.bold = True
p_eye.font.color.rgb = BORDEAUX
p_eye.space_after = Pt(2)

# Card Title
p_title = tf_fr.add_paragraph()
p_title.text = "Functional Requirements (FR)"
p_title.font.name = "Arial"
p_title.font.size = Pt(14)
p_title.font.bold = True
p_title.font.color.rgb = BORDEAUX_DARK
p_title.space_after = Pt(7)

# Bullets
fr_bullets = [
    ("FR1 — Multi-Modal 4D Navigation:", " Real-time toggle across synchronized 2D map, 3D WebGL point cloud, and 4D temporal playback."),
    ("FR2 — Automated Fault Plane Fitting:", " Mathematical extraction (RANSAC + SVD) computing Strike, Dip, and planarity ratio."),
    ("FR3 — Multi-Layer Geophysical Fusion:", " Volumetric 3D tomography (Vp/Vs), GNSS surface uplift deformation, and seismic stations."),
    ("FR4 — Dual Catalogue Ingestion:", " Automated queries to the INGV GOSSIP REST API and local offline CSV parsing engine.")
]

for tag, desc in fr_bullets:
    p = tf_fr.add_paragraph()
    r_tag = p.add_run()
    r_tag.text = f"• {tag}"
    r_tag.font.name = "Arial"
    r_tag.font.size = Pt(9.3)
    r_tag.font.bold = True
    r_tag.font.color.rgb = BORDEAUX_DARK
    
    r_desc = p.add_run()
    r_desc.text = desc
    r_desc.font.name = "Arial"
    r_desc.font.size = Pt(8.8)
    r_desc.font.color.rgb = TEXT_BODY
    p.space_after = Pt(4.5)


# ----------------- TOP-RIGHT CARD: Non-Functional Requirements (NFR) -----------------
# Seamless Rounded Rectangle with soft blue/ice tint (matching Slide 7 Card 2 & Slide 12 Card 1)
RIGHT_X = LEFT_MARGIN + CARD_W + GAP_X
nfr_card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, RIGHT_X, TOP_Y, CARD_W, TOP_H)
nfr_card.fill.solid()
nfr_card.fill.fore_color.rgb = BLUE_TINT_BG
nfr_card.line.color.rgb = BLUE_TINT_LN
nfr_card.line.width = Pt(1.2)

# Text inside NFR Card
nfr_text = slide4.shapes.add_textbox(RIGHT_X + Inches(0.22), TOP_Y + Inches(0.16), CARD_W - Inches(0.44), TOP_H - Inches(0.32))
tf_nfr = nfr_text.text_frame
tf_nfr.word_wrap = True
tf_nfr.margin_left = tf_nfr.margin_right = tf_nfr.margin_top = tf_nfr.margin_bottom = 0

# Eyebrow tag
p_eye = tf_nfr.paragraphs[0]
p_eye.text = "PERFORMANCE & SCIENTIFIC ROBUSTNESS"
p_eye.font.name = "Arial"
p_eye.font.size = Pt(9.2)
p_eye.font.bold = True
p_eye.font.color.rgb = NAVY_CATEGORY
p_eye.space_after = Pt(2)

# Card Title
p_title = tf_nfr.add_paragraph()
p_title.text = "Non-Functional Requirements (NFR)"
p_title.font.name = "Arial"
p_title.font.size = Pt(14)
p_title.font.bold = True
p_title.font.color.rgb = NAVY_CATEGORY
p_title.space_after = Pt(7)

# Bullets
nfr_bullets = [
    ("NFR1 — Real-Time Rendering:", " Sustained 50–60 FPS with up to 50,000 seismic hypocentres via WebGL GPU instancing."),
    ("NFR2 — Sub-Second Latency:", " Interactive clustering recomputation under 500 ms for instantaneous parameter tuning."),
    ("NFR3 — Bounded Footprint:", " Native cross-platform desktop wrapper (pywebview) strictly under 500 MB RAM consumption."),
    ("NFR4 — Robust Sanitisation:", " Automatic exclusion of zero-depth / corrupted events to avoid planar fitting distortions.")
]

for tag, desc in nfr_bullets:
    p = tf_nfr.add_paragraph()
    r_tag = p.add_run()
    r_tag.text = f"• {tag}"
    r_tag.font.name = "Arial"
    r_tag.font.size = Pt(9.3)
    r_tag.font.bold = True
    r_tag.font.color.rgb = NAVY_CATEGORY
    
    r_desc = p.add_run()
    r_desc.text = desc
    r_desc.font.name = "Arial"
    r_desc.font.size = Pt(8.8)
    r_desc.font.color.rgb = TEXT_BODY
    p.space_after = Pt(4.5)


# =========================================================================
# BOTTOM ROW: Minimalist Internship Timeline (Clean, Airy & Elegant)
# =========================================================================
BOTTOM_Y = 4120000
BOTTOM_H = 2440000

# Base Container (Soft Warm Ivory Rounded Rectangle matching Slide 12 bottom card)
timeline_card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, LEFT_MARGIN, BOTTOM_Y, TOTAL_W, BOTTOM_H)
timeline_card.fill.solid()
timeline_card.fill.fore_color.rgb = BEIGE_TINT_BG
timeline_card.line.color.rgb = BEIGE_TINT_LN
timeline_card.line.width = Pt(1.2)

# Header inside timeline card (Eyebrow + Title inline)
t_header = slide4.shapes.add_textbox(LEFT_MARGIN + Inches(0.24), BOTTOM_Y + Inches(0.14), TOTAL_W - Inches(0.48), Inches(0.35))
tf_th = t_header.text_frame
tf_th.word_wrap = True
tf_th.margin_left = tf_th.margin_right = tf_th.margin_top = tf_th.margin_bottom = 0
p = tf_th.paragraphs[0]

r_tag = p.add_run()
r_tag.text = "INTERNSHIP ROADMAP  •  "
r_tag.font.name = "Arial"
r_tag.font.size = Pt(9.2)
r_tag.font.bold = True
r_tag.font.color.rgb = BORDEAUX

r_th = p.add_run()
r_th.text = "Iterative Development Timeline"
r_th.font.name = "Arial"
r_th.font.size = Pt(12.5)
r_th.font.bold = True
r_th.font.color.rgb = TEXT_HEADING

r_th_sub = p.add_run()
r_th_sub.text = "   (6 progressive milestones shaped by direct feedback from INGV & Roma Tre)"
r_th_sub.font.name = "Arial"
r_th_sub.font.size = Pt(8.8)
r_th_sub.font.italic = True
r_th_sub.font.color.rgb = TEXT_MUTED

# Minimalist Timeline Metrics
phases = [
    ("Phase 1", "3D Prototype & CSV", "Hypocentre 3D cloud & local CSV parser engine.", "Three.js • CSV"),
    ("Phase 2", "K-Means & Bridge", "Spatial clustering & Python-WebView bridge.", "pywebview • IPC"),
    ("Phase 3", "UI Redesign & API", "Analytical dashboard & INGV GOSSIP API query.", "Leaflet • REST"),
    ("Phase 4", "DBSCAN & Cylinders", "Knee tuning, 4D cylinder slicing & locking.", "k-NN Knee • 4D"),
    ("Phase 5", "Geophysical Layers", "Vp/Vs tomography, GNSS surface & stations.", "Mesh • Tomography"),
    ("Phase 6", "Streaming & Async", "Float32 streaming (-80%) & worker pool.", "Binary • Workers")
]

num_phases = len(phases)
col_avail_w = TOTAL_W - Inches(0.48)
col_w = col_avail_w / num_phases

# Minimal Timeline Horizontal Axis
bar_y = BOTTOM_Y + Inches(0.70)
bar_line = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT_MARGIN + Inches(0.6), bar_y, TOTAL_W - Inches(1.2), Pt(2))
bar_line.fill.solid()
bar_line.fill.fore_color.rgb = TIMELINE_LINE_CLR
bar_line.line.fill.background()

# Nodes and Clean Text (NO heavy boxed sub-cards)
for i, (p_tag, p_title, p_desc, p_stack) in enumerate(phases):
    col_cx = LEFT_MARGIN + Inches(0.24) + i * col_w + col_w / 2
    
    # Milestone Circle Badge (Clean numbered circle sitting directly on the line)
    node_d = Inches(0.26)
    node = slide4.shapes.add_shape(MSO_SHAPE.OVAL, col_cx - node_d / 2, bar_y - node_d / 2, node_d, node_d)
    node.fill.solid()
    node.fill.fore_color.rgb = BORDEAUX_DARK if i == 5 else BORDEAUX
    node.line.color.rgb = RGBColor(255, 255, 255)
    node.line.width = Pt(2)
    
    tf_node = node.text_frame
    tf_node.margin_left = tf_node.margin_right = tf_node.margin_top = tf_node.margin_bottom = 0
    p_n = tf_node.paragraphs[0]
    p_n.alignment = PP_ALIGN.CENTER
    p_n.text = str(i + 1)
    p_n.font.name = "Arial"
    p_n.font.size = Pt(8.5)
    p_n.font.bold = True
    p_n.font.color.rgb = RGBColor(255, 255, 255)
    
    # Text container directly below node
    text_x = LEFT_MARGIN + Inches(0.24) + i * col_w + Inches(0.06)
    text_w = col_w - Inches(0.12)
    text_y = bar_y + Inches(0.20)
    text_h = BOTTOM_H - Inches(1.0)
    
    tbox = slide4.shapes.add_textbox(text_x, text_y, text_w, text_h)
    tf_t = tbox.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    
    # Phase Pill Label (e.g. PHASE 1)
    p_lbl = tf_t.paragraphs[0]
    p_lbl.text = p_tag.upper()
    p_lbl.font.name = "Arial"
    p_lbl.font.size = Pt(8.2)
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = BORDEAUX
    p_lbl.space_after = Pt(2)
    
    # Phase Title
    p_ti = tf_t.add_paragraph()
    p_ti.text = p_title
    p_ti.font.name = "Arial"
    p_ti.font.size = Pt(9.2)
    p_ti.font.bold = True
    p_ti.font.color.rgb = TEXT_HEADING
    p_ti.space_after = Pt(4)
    
    # Phase Description
    p_de = tf_t.add_paragraph()
    p_de.text = p_desc
    p_de.font.name = "Arial"
    p_de.font.size = Pt(8.2)
    p_de.font.color.rgb = TEXT_BODY
    p_de.space_after = Pt(6)
    
    # Focus tech pill text
    p_st = tf_t.add_paragraph()
    r_st_l = p_st.add_run()
    r_st_l.text = "Tech: "
    r_st_l.font.name = "Arial"
    r_st_l.font.size = Pt(7.8)
    r_st_l.font.bold = True
    r_st_l.font.color.rgb = BORDEAUX
    
    r_st_v = p_st.add_run()
    r_st_v.text = p_stack
    r_st_v.font.name = "Arial"
    r_st_v.font.size = Pt(7.8)
    r_st_v.font.color.rgb = TEXT_MUTED

# Save presentation
prs.save(pptx_path)
print("Saved minimalist Slide 4 to workspace pptx successfully!")

# Export slide 4 to PNG for visual inspection
try:
    ppt = win32com.client.Dispatch("PowerPoint.Application")
    pres = ppt.Presentations.Open(pptx_path, WithWindow=False)
    out_dir = os.path.join(script_dir, "exported_slides")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "slide_04.png")
    pres.Slides.Item(4).Export(out_path, "PNG", 1920, 1080)
    pres.Close()
    ppt.Quit()
    print(f"Exported minimalist slide 4 to {out_path}")
except Exception as e:
    print(f"PowerPoint export note: {e}")
