import pptx
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt, Emu
import os
import win32com.client

script_dir = os.path.dirname(os.path.abspath(__file__))
pptx_path = os.path.join(script_dir, "Seismic Viewer - Thesis Defense.pptx")
logo_path = os.path.join(script_dir, "sapienza_logo.png")
export_dir = os.path.join(script_dir, "exported_slides")

# Palette Definition
BORDEAUX = RGBColor(130, 36, 51)       # #822433
BORDEAUX_DARK = RGBColor(95, 20, 35)   # #5F1423
BORDEAUX_TINT = RGBColor(246, 237, 239)# #F6EDEF (soft bordeaux card highlight)
BEIGE_BG = RGBColor(250, 248, 242)     # #FAF8F2 (warm elegant background)
CARD_BG = RGBColor(255, 255, 255)      # #FFFFFF (clean white card)
CARD_LINE = RGBColor(228, 222, 212)    # #E4DED4 (warm beige border)
TEXT_HEADING = RGBColor(28, 24, 25)    # #1C1819 (main slide titles)
TEXT_BODY = RGBColor(45, 38, 40)       # #2D2628 (high-contrast body text)
TEXT_MUTED = RGBColor(115, 105, 103)   # #736967 (secondary/meta text)

prs = pptx.Presentation(pptx_path)

CARD_TITLES = {
    "Active Volcanic Unrest", "Planar Fault Structures", "INGV Operational Need",
    "GIS PLATFORMS", "SCIENTIFIC CODE", "OUR SOLUTION", "QGIS & ArcGIS", "ObsPy & GMT", "Seismic Viewer",
    "Functional Requirements (FR)", "Non-Functional Requirements (NFR)",
    "Target Users: Geophysicists, seismologists, and volcano monitoring operators.",
    "Constraints: Cross-platform support (Linux, macOS, Windows) and zero internet dependency.",
    "Zero-Latency Local Filtering", "Dual Layer Visualisation", "True Depth Scaling", "Aspect Ratio Correction",
    "Accumulation Mode", "Sliding Window", "Discrete Block Mode", "Interactive Timeline Histogram & Energy Meter",
    "Vp/Vs Seismic Tomography", "GNSS Ground Uplift Columns",
    "DBSCAN", "MiniBatchKMeans", "Cylinder Grid", "Hierarchical Sub-Clustering & Lock Mechanism",
    "Step 1: RANSAC Inlier Extraction", "Step 2: SVD & Geological Angles",
    "Why Cylinder Partitioning?", "Global Label Matching",
    "STAGE 1 & 2", "CRITICAL DECISION", "STAGE 3 & 4", "STAGE 5",
    "Multi-Source Loading", "Sentinel Depth Handling", "Physical Bounds Check", "Geospatial Projection",
    "Float32 Contiguous Binary Protocol (/api/data/{id}/bin)",
    "Synthetic Monte Carlo Validation (500 runs)", "Real Catalogue Testing (Campi Flegrei)",
    "Serialization & Memory Footprint", "Algorithmic Throughput",
    "Validated Contributions", "Future Extensions",
    "Candidate", "Academic Supervisor", "Co-Supervisor"
}

for idx, slide in enumerate(prs.slides, 1):
    # Ensure beige background
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = BEIGE_BG

    for s in slide.shapes:
        if not s.has_text_frame:
            continue
        
        for p in s.text_frame.paragraphs:
            txt = p.text.strip()
            if not txt:
                continue
            
            # Determine target color for all runs in paragraph
            is_super = (s.top < 650000 and s.height < 250000 and txt.isupper() and len(txt) > 4)
            is_title = (s.top < 1250000 and s.height >= 300000)
            is_metric = any(m in txt for m in ["< 0.51°", "< 0.93°", "-80% Size", "< 0.01 ms"])
            is_card_title = (txt in CARD_TITLES) or (s.height < 280000 and len(txt) < 45 and not txt.startswith("•") and not txt.startswith("-"))
            is_person = txt in ["Alessio Curcio", "Prof. Danilo Avola", "Dr. Daniele Bailo"]

            # Choose color
            if idx == 1:
                if txt == "Seismic Viewer":
                    color = BORDEAUX
                    bold = True
                elif "Design, Implementation" in txt:
                    color = TEXT_HEADING
                    bold = False
                elif is_person:
                    color = TEXT_HEADING
                    bold = True
                elif txt in ["Candidate", "Academic Supervisor", "Co-Supervisor"]:
                    color = BORDEAUX
                    bold = True
                elif "SAPIENZA" in txt:
                    color = BORDEAUX
                    bold = True
                else:
                    color = TEXT_BODY
                    bold = False
            elif is_super:
                color = BORDEAUX
                bold = True
            elif is_title:
                color = TEXT_HEADING
                bold = True
            elif is_metric:
                color = BORDEAUX
                bold = True
            elif is_card_title:
                if txt in ["STAGE 1 & 2", "CRITICAL DECISION", "STAGE 3 & 4", "STAGE 5"]:
                    color = BORDEAUX
                    bold = True
                elif txt in ["GIS PLATFORMS", "SCIENTIFIC CODE"]:
                    color = TEXT_MUTED
                    bold = True
                else:
                    color = BORDEAUX
                    bold = True
            elif is_person:
                color = TEXT_HEADING
                bold = True
            else:
                color = TEXT_BODY
                bold = False

            # Set color on paragraph and on EVERY RUN
            p.font.color.rgb = color
            p.font.bold = bold
            for r in p.runs:
                r.font.color.rgb = color
                r.font.bold = bold

prs.save(pptx_path)
print("Updated all text run colors successfully!")

# Re-export to PNG
ppt = win32com.client.Dispatch("PowerPoint.Application")
try:
    pres = ppt.Presentations.Open(pptx_path, WithWindow=False)
    for i, slide in enumerate(pres.Slides, 1):
        out_path = os.path.join(export_dir, f"slide_{i:02d}.png")
        slide.Export(out_path, "PNG", 1920, 1080)
    pres.Close()
finally:
    ppt.Quit()
print("Re-exported slides successfully!")
