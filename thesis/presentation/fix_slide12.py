import pptx
from pptx.util import Pt
from pptx.dml.color import RGBColor

BORDEAUX      = RGBColor(130, 36, 51)    # #822433
SLATE_NAVY    = RGBColor(32, 64, 96)     # #204060
SAGE_TEAL     = RGBColor(35, 107, 94)    # #236B5E
TERRACOTTA    = RGBColor(184, 88, 40)    # #B85828
GOLD_AMBER    = RGBColor(246, 173, 85)   # #F6AD55
CHARCOAL_DARK = RGBColor(28, 24, 25)     # #1C1819
DARK_NAVY     = RGBColor(30, 41, 59)     # #1E293B
WHITE         = RGBColor(255, 255, 255)
LIGHT_SLATE   = RGBColor(203, 213, 225)

TINT_NAVY     = RGBColor(242, 246, 250)  # #F2F6FA
TINT_AMBER    = RGBColor(253, 248, 242)  # #FDF8F2
TINT_SAGE     = RGBColor(242, 248, 246)  # #F2F8F6
TINT_BORDEAUX = RGBColor(251, 244, 243)  # #FBF4F3

prs = pptx.Presentation('thesis/presentation/Seismic Viewer - Thesis Defense.pptx')
s12 = prs.slides[11]

# Set coordinates for all shapes in Slide 12
for sh in s12.shapes:
    # 4 Top Cards
    if sh.name == 'Google Shape;301;p24':
        sh.left = 609600; sh.top = 1950000; sh.width = 2628900; sh.height = 2050000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_NAVY
        sh.line.color.rgb = RGBColor(208, 220, 229)
    elif sh.name == 'Google Shape;302;p24':
        sh.left = 3390900; sh.top = 1950000; sh.width = 2628900; sh.height = 2050000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_AMBER
        sh.line.color.rgb = RGBColor(235, 215, 197)
    elif sh.name == 'Google Shape;303;p24':
        sh.left = 6172200; sh.top = 1950000; sh.width = 2628900; sh.height = 2050000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_SAGE
        sh.line.color.rgb = RGBColor(205, 224, 219)
    elif sh.name == 'Google Shape;304;p24':
        sh.left = 8953500; sh.top = 1950000; sh.width = 2628900; sh.height = 2050000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_BORDEAUX
        sh.line.color.rgb = RGBColor(232, 213, 216)

    # Categories
    elif sh.name == 'Google Shape;307;p24':
        sh.left = 781050; sh.top = 2100000; sh.width = 2381250; sh.height = 200000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(10.5); p.font.bold = True; p.font.color.rgb = SLATE_NAVY
            for r in p.runs: r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = SLATE_NAVY
    elif sh.name == 'Google Shape;308;p24':
        sh.left = 3562350; sh.top = 2100000; sh.width = 2381250; sh.height = 200000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(10.5); p.font.bold = True; p.font.color.rgb = TERRACOTTA
            for r in p.runs: r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = TERRACOTTA
    elif sh.name == 'Google Shape;309;p24':
        sh.left = 6343650; sh.top = 2100000; sh.width = 2381250; sh.height = 200000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(10.5); p.font.bold = True; p.font.color.rgb = SAGE_TEAL
            for r in p.runs: r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = SAGE_TEAL
    elif sh.name == 'Google Shape;310;p24':
        sh.left = 9124950; sh.top = 2100000; sh.width = 2381250; sh.height = 200000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(10.5); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = BORDEAUX

    # Card Titles
    elif sh.name == 'Google Shape;312;p24':
        sh.left = 781050; sh.top = 2320000; sh.width = 2381250; sh.height = 250000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(15.5); p.font.bold = True; p.font.color.rgb = SLATE_NAVY
            for r in p.runs: r.font.size = Pt(15.5); r.font.bold = True; r.font.color.rgb = SLATE_NAVY
    elif sh.name == 'Google Shape;314;p24':
        sh.left = 3562350; sh.top = 2320000; sh.width = 2381250; sh.height = 250000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(15.5); p.font.bold = True; p.font.color.rgb = TERRACOTTA
            for r in p.runs: r.font.size = Pt(15.5); r.font.bold = True; r.font.color.rgb = TERRACOTTA
    elif sh.name == 'Google Shape;316;p24':
        sh.left = 6343650; sh.top = 2320000; sh.width = 2381250; sh.height = 250000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(15.5); p.font.bold = True; p.font.color.rgb = SAGE_TEAL
            for r in p.runs: r.font.size = Pt(15.5); r.font.bold = True; r.font.color.rgb = SAGE_TEAL
    elif sh.name == 'Google Shape;318;p24':
        sh.left = 9124950; sh.top = 2320000; sh.width = 2381250; sh.height = 250000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(15.5); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(15.5); r.font.bold = True; r.font.color.rgb = BORDEAUX

    # Card Body Texts
    elif sh.name == 'Google Shape;313;p24':
        sh.left = 781050; sh.top = 2700000; sh.width = 2362200; sh.height = 1150000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;315;p24':
        sh.left = 3562350; sh.top = 2700000; sh.width = 2362200; sh.height = 1150000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;317;p24':
        sh.left = 6343650; sh.top = 2700000; sh.width = 2362200; sh.height = 1150000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;319;p24':
        sh.left = 9124950; sh.top = 2700000; sh.width = 2362200; sh.height = 1150000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK

    # Bottom Hero Card Container
    elif sh.name == 'Google Shape;300;p24':
        sh.left = 609600; sh.top = 4250000; sh.width = 10972800; sh.height = 1050000
        sh.fill.solid(); sh.fill.fore_color.rgb = DARK_NAVY
        sh.line.fill.background()

    # Bottom Hero Card Text
    elif sh.name == 'Google Shape;320;p24': # Float32 Protocol Title
        sh.left = 885825; sh.top = 4420000; sh.width = 8000000; sh.height = 250000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16.5); p.font.bold = True; p.font.color.rgb = WHITE
            for r in p.runs: r.font.size = Pt(16.5); r.font.bold = True; r.font.color.rgb = WHITE
    elif sh.name == 'Google Shape;321;p24': # Subtitle
        sh.left = 885825; sh.top = 4780000; sh.width = 8000000; sh.height = 350000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = LIGHT_SLATE
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = LIGHT_SLATE

    # Badges
    elif sh.name == 'Google Shape;311;p24': # -80% Size
        sh.left = 9200000; sh.top = 4400000; sh.width = 2200000; sh.height = 420000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(26); p.font.bold = True; p.font.color.rgb = GOLD_AMBER
            for r in p.runs: r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = GOLD_AMBER
    elif sh.name == 'Google Shape;322;p24': # < 0.01 ms decode
        sh.left = 9200000; sh.top = 4850000; sh.width = 2200000; sh.height = 280000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = WHITE
            for r in p.runs: r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = WHITE

prs.save('thesis/presentation/Seismic Viewer - Thesis Defense.pptx')
print('Slide 12 fixed successfully!')
