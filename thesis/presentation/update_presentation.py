import pptx
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

BORDEAUX      = RGBColor(130, 36, 51)    # #822433
SLATE_NAVY    = RGBColor(32, 64, 96)     # #204060
SAGE_TEAL     = RGBColor(35, 107, 94)    # #236B5E
TERRACOTTA    = RGBColor(184, 88, 40)    # #B85828
GOLD_AMBER    = RGBColor(185, 125, 30)   # #B97D1E
CHARCOAL_DARK = RGBColor(28, 24, 25)     # #1C1819
TEXT_MUTED    = RGBColor(90, 80, 82)     # #5A5052
DARK_NAVY     = RGBColor(30, 41, 59)     # #1E293B
WHITE         = RGBColor(255, 255, 255)

TINT_BORDEAUX = RGBColor(251, 244, 243)  # #FBF4F3
TINT_NAVY     = RGBColor(242, 246, 250)  # #F2F6FA
TINT_SAGE     = RGBColor(242, 248, 246)  # #F2F8F6
TINT_AMBER    = RGBColor(253, 248, 242)  # #FDF8F2

prs = pptx.Presentation('thesis/presentation/Seismic Viewer - Thesis Defense.pptx')

# -------------------------------------------------------------
# SLIDE 2: Campi Flegrei (Raise cards, increase font size, color accents)
# -------------------------------------------------------------
s2 = prs.slides[1]
for sh in s2.shapes:
    if sh.name == 'Google Shape;102;p14':
        sh.top = 2200000; sh.height = 2850000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_AMBER
        sh.line.color.rgb = RGBColor(235, 215, 197)
    elif sh.name == 'Google Shape;103;p14':
        sh.top = 2200000; sh.height = 2850000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_NAVY
        sh.line.color.rgb = RGBColor(208, 220, 229)
    elif sh.name == 'Google Shape;104;p14':
        sh.top = 2200000; sh.height = 2850000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_BORDEAUX
        sh.line.color.rgb = RGBColor(232, 213, 216)
    elif sh.name in ['Google Shape;105;p14', 'Google Shape;106;p14', 'Google Shape;107;p14']:
        sh.top = 2470000
    elif sh.name in ['Google Shape;108;p14', 'Google Shape;109;p14', 'Google Shape;110;p14']:
        sh.top = 2580000
    elif sh.name == 'Google Shape;113;p14': # Card 1 title
        sh.top = 3050000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = TERRACOTTA
            for r in p.runs: r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = TERRACOTTA
    elif sh.name == 'Google Shape;114;p14': # Card 1 body
        sh.top = 3450000; sh.height = 1400000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13.5); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13.5); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;115;p14': # Card 2 title
        sh.top = 3050000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = SLATE_NAVY
            for r in p.runs: r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = SLATE_NAVY
    elif sh.name == 'Google Shape;116;p14': # Card 2 body
        sh.top = 3450000; sh.height = 1400000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13.5); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13.5); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;117;p14': # Card 3 title
        sh.top = 3050000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = BORDEAUX
    elif sh.name == 'Google Shape;118;p14': # Card 3 body
        sh.top = 3450000; sh.height = 1400000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13.5); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13.5); r.font.color.rgb = CHARCOAL_DARK

# -------------------------------------------------------------
# SLIDE 3: Tooling Comparison (Raise cards, increase font size, color accents)
# -------------------------------------------------------------
s3 = prs.slides[2]
for sh in s3.shapes:
    if sh.name == 'Google Shape;123;p15': # QGIS
        sh.top = 2200000; sh.height = 2850000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_NAVY
        sh.line.color.rgb = RGBColor(208, 220, 229)
    elif sh.name == 'Google Shape;124;p15': # ObsPy
        sh.top = 2200000; sh.height = 2850000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_NAVY
        sh.line.color.rgb = RGBColor(208, 220, 229)
    elif sh.name == 'Google Shape;125;p15': # Seismic Viewer (Hero card)
        sh.top = 2200000; sh.height = 2850000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_BORDEAUX
        sh.line.color.rgb = BORDEAUX
        sh.line.width = Pt(1.5)
    elif sh.name in ['Google Shape;128;p15', 'Google Shape;130;p15']:
        sh.top = 2400000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(10.5); p.font.bold = True; p.font.color.rgb = SLATE_NAVY
            for r in p.runs: r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = SLATE_NAVY
    elif sh.name == 'Google Shape;132;p15': # OUR SOLUTION
        sh.top = 2400000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(10.5); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = BORDEAUX
    elif sh.name in ['Google Shape;129;p15', 'Google Shape;131;p15']:
        sh.top = 2620000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;133;p15': # Seismic Viewer title
        sh.top = 2620000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = BORDEAUX
    elif sh.name in [f'Google Shape;{i};p15' for i in range(134, 143)]:
        sh.top = sh.top - 500000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK

# -------------------------------------------------------------
# SLIDE 5: 2D Mode (Bump small fonts to 13.5pt, add headers, coherent colors)
# -------------------------------------------------------------
s5 = prs.slides[4]
for sh in s5.shapes:
    if sh.name == 'Google Shape;173;p17': # Card 0
        sh.top = 1890013; sh.height = 1180000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_NAVY
        sh.line.color.rgb = RGBColor(208, 220, 229)
    elif sh.name == 'Google Shape;174;p17': # Card 1
        sh.top = 3190013; sh.height = 1180000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_BORDEAUX
        sh.line.color.rgb = RGBColor(232, 213, 216)
    elif sh.name == 'Google Shape;175;p17': # Card 2
        sh.top = 4490013; sh.height = 1180000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_SAGE
        sh.line.color.rgb = RGBColor(205, 224, 219)
    elif sh.name == 'Google Shape;180;p17': # Card 0 header
        sh.top = 1990000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = SLATE_NAVY
            for r in p.runs: r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = SLATE_NAVY
    elif sh.name == 'Google Shape;181;p17': # Card 0 body
        sh.top = 2280000; sh.height = 680000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13.5); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13.5); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;182;p17': # Card 1 header
        sh.top = 3290000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = BORDEAUX
    elif sh.name == 'Google Shape;183;p17': # Card 1 body
        sh.top = 3580000; sh.height = 680000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13.5); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13.5); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;184;p17': # Card 2 text -> Format as header + body
        sh.top = 4590000; sh.height = 980000
        sh.text_frame.clear()
        p1 = sh.text_frame.paragraphs[0]
        p1.text = 'Spatio-Temporal Swarm Isolation'
        p1.font.size = Pt(16); p1.font.bold = True; p1.font.color.rgb = SAGE_TEAL
        p2 = sh.text_frame.add_paragraph()
        p2.text = 'Immediate epicentral clustering identification across Pozzuoli and Solfatara crater zones.'
        p2.font.size = Pt(13.5); p2.font.color.rgb = CHARCOAL_DARK

# -------------------------------------------------------------
# SLIDE 6: 3D Mode (Add header to 3rd card, bump font sizes)
# -------------------------------------------------------------
s6 = prs.slides[5]
for sh in s6.shapes:
    if sh.name == 'Google Shape;189;p18':
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_BORDEAUX
        sh.line.color.rgb = RGBColor(232, 213, 216)
    elif sh.name == 'Google Shape;190;p18':
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_NAVY
        sh.line.color.rgb = RGBColor(208, 220, 229)
    elif sh.name == 'Google Shape;191;p18':
        sh.height = 1000000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_SAGE
        sh.line.color.rgb = RGBColor(205, 224, 219)
    elif sh.name == 'Google Shape;196;p18':
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = BORDEAUX
    elif sh.name == 'Google Shape;197;p18':
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13.5); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13.5); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;198;p18':
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = SLATE_NAVY
            for r in p.runs: r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = SLATE_NAVY
    elif sh.name == 'Google Shape;199;p18':
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13.5); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13.5); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;200;p18': # Card 3 text -> Header + body
        sh.height = 800000
        sh.text_frame.clear()
        p1 = sh.text_frame.paragraphs[0]
        p1.text = 'Stateful Camera Retention'
        p1.font.size = Pt(16); p1.font.bold = True; p1.font.color.rgb = SAGE_TEAL
        p2 = sh.text_frame.add_paragraph()
        p2.text = "Camera viewpoint preserved across filter updates via Plotly uirevision='constant'."
        p2.font.size = Pt(13.5); p2.font.color.rgb = CHARCOAL_DARK

# -------------------------------------------------------------
# SLIDE 12: Data Engineering (High-tech Slate Navy bottom card)
# -------------------------------------------------------------
s12 = prs.slides[11]
for sh in s12.shapes:
    if sh.name == 'Google Shape;304;p24': # Bottom hero card
        sh.fill.solid(); sh.fill.fore_color.rgb = DARK_NAVY
        sh.line.fill.background()
    elif sh.name == 'Google Shape;315;p24': # Title inside bottom hero card
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = RGBColor(255, 255, 255)
            for r in p.runs: r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
    elif sh.name == 'Google Shape;316;p24': # Subtitle inside bottom hero card
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = RGBColor(203, 213, 225)
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = RGBColor(203, 213, 225)
    elif sh.name == 'BadgeText_12_main': # -80% Size
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(26); p.font.bold = True; p.font.color.rgb = RGBColor(246, 173, 85) # Gold Amber
            for r in p.runs: r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = RGBColor(246, 173, 85)
    elif sh.name == 'BadgeText_12_sub': # < 0.01 ms decode
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = RGBColor(255, 255, 255)
            for r in p.runs: r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
    elif sh.name in ['Google Shape;300;p24', 'Google Shape;301;p24', 'Google Shape;302;p24', 'Google Shape;303;p24']:
        sh.top = 2150000; sh.height = 2100000
    elif sh.name == 'Google Shape;308;p24': # Stage 1 header
        sh.top = 2450000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(15.5); p.font.bold = True; p.font.color.rgb = SLATE_NAVY
            for r in p.runs: r.font.size = Pt(15.5); r.font.bold = True; r.font.color.rgb = SLATE_NAVY
    elif sh.name == 'Google Shape;310;p24': # Stage 2 header
        sh.top = 2450000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(15.5); p.font.bold = True; p.font.color.rgb = TERRACOTTA
            for r in p.runs: r.font.size = Pt(15.5); r.font.bold = True; r.font.color.rgb = TERRACOTTA
    elif sh.name == 'Google Shape;312;p24': # Stage 3 header
        sh.top = 2450000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(15.5); p.font.bold = True; p.font.color.rgb = SAGE_TEAL
            for r in p.runs: r.font.size = Pt(15.5); r.font.bold = True; r.font.color.rgb = SAGE_TEAL
    elif sh.name == 'Google Shape;314;p24': # Stage 4 header
        sh.top = 2450000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(15.5); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(15.5); r.font.bold = True; r.font.color.rgb = BORDEAUX
    elif sh.name in ['Google Shape;307;p24', 'Google Shape;309;p24', 'Google Shape;311;p24', 'Google Shape;313;p24']:
        sh.top = 2250000
    elif sh.name in [f'Google Shape;{i};p24' for i in [317, 318, 319, 320]]:
        sh.top = 2850000; sh.height = 1300000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK

# -------------------------------------------------------------
# SLIDE 13: Results & Benchmarks (Raise cards, bump font sizes, Sage Teal badges)
# -------------------------------------------------------------
s13 = prs.slides[12]
for sh in s13.shapes:
    if sh.name in ['Google Shape;327;p25', 'Google Shape;328;p25']:
        sh.top = 2150000; sh.height = 1500000
    elif sh.name in ['Google Shape;329;p25', 'Google Shape;330;p25']:
        sh.top = 3850000; sh.height = 1500000
    elif sh.name == 'Google Shape;333;p25': # Synthetic Monte Carlo
        sh.top = 2250000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = BORDEAUX
    elif sh.name == 'Google Shape;334;p25': # Monte Carlo body
        sh.top = 2600000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name in ['Google Shape;335;p25', 'Google Shape;337;p25']: # < 0.51° and < 0.93°
        sh.top = 3100000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = SAGE_TEAL
            for r in p.runs: r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = SAGE_TEAL
    elif sh.name in ['Google Shape;336;p25', 'Google Shape;338;p25']:
        sh.top = 3450000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = TEXT_MUTED
            for r in p.runs: r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = TEXT_MUTED
    elif sh.name == 'Google Shape;339;p25':
        sh.top = 2250000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = TERRACOTTA
            for r in p.runs: r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = TERRACOTTA
    elif sh.name == 'Google Shape;340;p25':
        sh.top = 2600000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;341;p25':
        sh.top = 3950000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = BORDEAUX
    elif sh.name == 'Google Shape;342;p25':
        sh.top = 4300000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;343;p25':
        sh.top = 3950000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = SLATE_NAVY
            for r in p.runs: r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = SLATE_NAVY
    elif sh.name == 'Google Shape;344;p25':
        sh.top = 4300000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK

# -------------------------------------------------------------
# SLIDE 14: Conclusions (Raise cards, Sage Teal & Slate Navy)
# -------------------------------------------------------------
s14 = prs.slides[13]
for sh in s14.shapes:
    if sh.name == 'Google Shape;354;p26': # Left card
        sh.top = 2200000; sh.height = 2700000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_SAGE
        sh.line.color.rgb = RGBColor(205, 224, 219)
    elif sh.name == 'Google Shape;355;p26': # Right card
        sh.top = 2200000; sh.height = 2700000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_NAVY
        sh.line.color.rgb = RGBColor(208, 220, 229)
    elif sh.name == 'Google Shape;356;p26':
        sh.top = 2470000
    elif sh.name == 'Google Shape;357;p26':
        sh.top = 2470000
    elif sh.name == 'Google Shape;364;p26':
        sh.top = 2450000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = SAGE_TEAL
            for r in p.runs: r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = SAGE_TEAL
    elif sh.name == 'Google Shape;365;p26':
        sh.top = 2850000; sh.height = 1800000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(14); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(14); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;366;p26':
        sh.top = 2450000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = SLATE_NAVY
            for r in p.runs: r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = SLATE_NAVY
    elif sh.name == 'Google Shape;367;p26':
        sh.top = 2850000; sh.height = 1800000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(14); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(14); r.font.color.rgb = CHARCOAL_DARK

prs.save('thesis/presentation/Seismic Viewer - Thesis Defense.pptx')
print('Script completed successfully!')
