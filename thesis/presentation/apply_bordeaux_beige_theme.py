import pptx
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt, Emu
from PIL import Image
import os
import shutil

# Paths
script_dir = os.path.dirname(os.path.abspath(__file__))
pptx_path = os.path.join(script_dir, "Seismic Viewer - Thesis Defense.pptx")
logo_path = os.path.join(script_dir, "sapienza_logo.png")

# Palette Definition
# Bordeaux Sapienza
BORDEAUX = RGBColor(130, 36, 51)       # #822433
BORDEAUX_DARK = RGBColor(100, 20, 35)  # #641423
BORDEAUX_TINT = RGBColor(246, 237, 239)# #F6EDEF (soft bordeaux card highlight)
BORDEAUX_LINE = RGBColor(180, 60, 78)  # #B43C4E

# Warm Academic Beige
BEIGE_BG = RGBColor(250, 248, 242)     # #FAF8F2 (warm elegant background)
CARD_BG = RGBColor(255, 255, 255)      # #FFFFFF (clean white card)
CARD_LINE = RGBColor(228, 222, 212)    # #E4DED4 (warm beige border)
CARD_LINE_LIGHT = RGBColor(238, 234, 226)

# Text Colors
TEXT_TITLE = BORDEAUX                  # #822433 (titles & category superheaders)
TEXT_HEADING = RGBColor(28, 24, 25)    # #1C1819 (main slide titles)
TEXT_CARD_TITLE = BORDEAUX             # #822433 (card titles)
TEXT_BODY = RGBColor(48, 42, 44)       # #302A2C (readable high-contrast body)
TEXT_MUTED = RGBColor(115, 105, 103)   # #736967 (secondary/meta text)

print("Opening presentation...")
prs = pptx.Presentation(pptx_path)

# 1. Slide 2 Fix: Equalize Box Dimensions
slide2 = prs.slides[1]
# Shape 0: Box 0 (left=609600, top=2690812, w=3505200, h=2324100)
# Shape 1: Box 1 (left=4343400, top=2797968, w=3505200, h=2109787) -> Make same top & height!
box1 = slide2.shapes[1]
box1.top = 2690812
box1.height = 2324100

# Shift inner elements of Box 1 up by 107156 EMU (2797968 - 2690812)
dy = 107156
for s in [slide2.shapes[4], slide2.shapes[7], slide2.shapes[13], slide2.shapes[14]]:
    s.top = s.top - dy
slide2.shapes[14].height = 857250 # match height of other description boxes
print("Slide 2: Box 1 dimensions and positions equalized!")

# 2. Slide 5 Fix: Reduce Card Dimensions
slide5 = prs.slides[4]
# Card 0, 1, 2
# Currently w=4791967. Reduce to w=4400000.
# Card 0: top=2150000, h=980000
# Card 1: top=3250000, h=980000
# Card 2: top=4350000, h=800000
c0, c1, c2 = slide5.shapes[0], slide5.shapes[1], slide5.shapes[2]
c0.left = 609600; c0.top = 2150000; c0.width = 4400000; c0.height = 980000
c1.left = 609600; c1.top = 3250000; c1.width = 4400000; c1.height = 980000
c2.left = 609600; c2.top = 4350000; c2.width = 4400000; c2.height = 800000

# Inner text boxes for Card 0
# Title (shape 7)
slide5.shapes[7].left = 800000; slide5.shapes[7].top = 2280000; slide5.shapes[7].width = 4050000; slide5.shapes[7].height = 220000
# Desc (shape 8)
slide5.shapes[8].left = 800000; slide5.shapes[8].top = 2520000; slide5.shapes[8].width = 4050000; slide5.shapes[8].height = 480000

# Inner text boxes for Card 1
# Title (shape 9)
slide5.shapes[9].left = 800000; slide5.shapes[9].top = 3380000; slide5.shapes[9].width = 4050000; slide5.shapes[9].height = 220000
# Desc (shape 10)
slide5.shapes[10].left = 800000; slide5.shapes[10].top = 3620000; slide5.shapes[10].width = 4050000; slide5.shapes[10].height = 480000

# Inner text box for Card 2 (shape 11)
slide5.shapes[11].left = 800000; slide5.shapes[11].top = 4480000; slide5.shapes[11].width = 4050000; slide5.shapes[11].height = 540000
print("Slide 5: Card dimensions reduced and re-padded!")

# 3. Slide 12 Fix: Text Overflow on -80% size badge
slide12 = prs.slides[11]
# Container shape 0: left=609600, top=4209008, w=10972800, h=914400
# Title (shape 20): left=885825, top=4350000, w=8000000, h=220000
slide12.shapes[20].left = 885825; slide12.shapes[20].top = 4360000; slide12.shapes[20].width = 8000000; slide12.shapes[20].height = 220000
# Subtitle (shape 21): left=885825, top=4620000, w=8000000, h=260000
slide12.shapes[21].left = 885825; slide12.shapes[21].top = 4620000; slide12.shapes[21].width = 8000000; slide12.shapes[21].height = 260000

# Stat badge (-80% Size, shape 11) & Subtitle (< 0.01 ms decode time, shape 22)
# Place them neatly inside a dedicated right stat area
slide12.shapes[11].left = 9150000; slide12.shapes[11].top = 4340000; slide12.shapes[11].width = 2250000; slide12.shapes[11].height = 340000
slide12.shapes[22].left = 9150000; slide12.shapes[22].top = 4690000; slide12.shapes[22].width = 2250000; slide12.shapes[22].height = 250000

# Center text in shape 11 and 22
for s in [slide12.shapes[11], slide12.shapes[22]]:
    if s.has_text_frame:
        for p in s.text_frame.paragraphs:
            p.alignment = pptx.enum.text.PP_ALIGN.CENTER
print("Slide 12: Fixed badge alignment and text overflow!")

# 4. Apply Beige Background & Style all Slides
for idx, slide in enumerate(prs.slides, 1):
    # Set Slide Background to Warm Beige
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BEIGE_BG

    # Style shapes
    for s in slide.shapes:
        # Check if shape is a card / container
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.AUTO_SHAPE:
            # Distinguish thin divider lines from cards
            if s.height < 20000 or s.width < 20000:
                # Divider line -> Bordeaux accent line
                try:
                    s.fill.solid()
                    s.fill.fore_color.rgb = BORDEAUX
                except Exception:
                    pass
            elif idx == 3 and s.name == "Google Shape;125;p15":
                # Slide 3 "Our Solution" card -> Highlighted with soft Bordeaux tint
                try:
                    s.fill.solid()
                    s.fill.fore_color.rgb = BORDEAUX_TINT
                    s.line.color.rgb = BORDEAUX
                    s.line.width = Emu(19050)
                except Exception:
                    pass
            elif idx == 12 and s.name == "Google Shape;300;p24":
                # Slide 12 Float32 Binary Container -> Highlighted with soft Bordeaux tint
                try:
                    s.fill.solid()
                    s.fill.fore_color.rgb = BORDEAUX_TINT
                    s.line.color.rgb = BORDEAUX
                    s.line.width = Emu(19050)
                except Exception:
                    pass
            else:
                # Regular cards -> Clean white with warm beige border
                try:
                    s.fill.solid()
                    s.fill.fore_color.rgb = CARD_BG
                    s.line.color.rgb = CARD_LINE
                    s.line.width = Emu(12700)
                except Exception:
                    pass

        # Style Text Elements
        if s.has_text_frame:
            for p_idx, p in enumerate(s.text_frame.paragraphs):
                # Analyze paragraph content to determine semantic role
                txt = p.text.strip()
                if not txt:
                    continue
                
                # Check for category super-header (all caps at top of slide)
                if s.top < 650000 and s.height < 250000 and txt.isupper() and len(txt) > 5:
                    p.font.color.rgb = BORDEAUX
                    p.font.bold = True
                # Main slide title (e.g. "The Campi Flegrei Volcanic Challenge")
                elif s.top < 1200000 and s.height >= 300000:
                    p.font.color.rgb = TEXT_HEADING
                    p.font.bold = True
                # Card Titles / Headings
                elif (s.height < 300000 and len(txt) < 50 and not txt.startswith("•")) or txt in [
                    "Active Volcanic Unrest", "Planar Fault Structures", "INGV Operational Need",
                    "GIS PLATFORMS", "SCIENTIFIC CODE", "OUR SOLUTION", "QGIS & ArcGIS", "ObsPy & GMT", "Seismic Viewer",
                    "Functional Requirements (FR)", "Non-Functional Requirements (NFR)",
                    "Target Users: Geophysicists, seismologists, and volcano monitoring operators.",
                    "Constraints: Cross-platform support (Linux, macOS, Windows) and zero internet dependency.",
                    "Zero-Latency Local Filtering", "Dual Layer Visualisation", "True Depth Scaling", "Aspect Ratio Correction",
                    "Accumulation Mode", "Sliding Window", "Discrete Block Mode", "Interactive Timeline Histogram & Energy Meter",
                    "Vp/Vs Seismic Tomography", "GNSS Ground Uplift Columns",
                    "DBSCAN", "MiniBatchKMeans", "Cylinder Grid", "Hierarchical Sub-Clustering & Lock Mechanism",
                    "Step 1: RANSAC Inlier Extraction", "Step 2: SVD & Geological Angles",
                    "Why Cylinder Partitioning?", "Global Label Matching",
                    "STAGE 1 & 2", "CRITICAL DECISION", "STAGE 3 & 4", "STAGE 5",
                    "Multi-Source Loading", "Sentinel Depth Handling", "Physical Bounds Check", "Geospatial Projection",
                    "Float32 Contiguous Binary Protocol (/api/data/{id}/bin)",
                    "Synthetic Monte Carlo Validation (500 runs)", "Real Catalogue Testing (Campi Flegrei)",
                    "Serialization & Memory Footprint", "Algorithmic Throughput",
                    "Validated Contributions", "Future Extensions",
                    "Candidate", "Academic Supervisor", "Co-Supervisor"
                ]:
                    if txt in ["GIS PLATFORMS", "SCIENTIFIC CODE", "STAGE 1 & 2", "CRITICAL DECISION", "STAGE 3 & 4", "STAGE 5"]:
                        p.font.color.rgb = TEXT_MUTED
                        p.font.bold = True
                    else:
                        p.font.color.rgb = TEXT_CARD_TITLE
                        p.font.bold = True
                # Key Metrics / Numbers (e.g. "< 0.51°", "< 0.93°", "-80% Size", "< 0.01 ms decode time")
                elif any(m in txt for m in ["< 0.51°", "< 0.93°", "-80% Size"]):
                    p.font.color.rgb = BORDEAUX
                    p.font.bold = True
                # Slide 1 Main Title
                elif idx == 1 and txt == "Seismic Viewer":
                    p.font.color.rgb = BORDEAUX
                    p.font.bold = True
                # Candidate name on Slide 1 & 14
                elif txt in ["Alessio Curcio", "Prof. Danilo Avola", "Dr. Daniele Bailo"]:
                    p.font.color.rgb = TEXT_HEADING
                    p.font.bold = True
                # Body text inside cards
                else:
                    p.font.color.rgb = TEXT_BODY

print("Beige and Bordeaux styling applied across all 14 slides!")

# 5. Add Sapienza Logo
# Add logo on Slide 1 (Title), Slide 14 (Conclusion), and discreetly on Slides 2-13
if os.path.exists(logo_path):
    print("Embedding Sapienza logo...")
    # Slide 1: Prominent header logo at top right
    slide1 = prs.slides[0]
    slide1.shapes.add_picture(logo_path, left=8800000, top=520000, width=2600000, height=807000)

    # Slide 14: Top right header logo
    slide14 = prs.slides[13]
    slide14.shapes.add_picture(logo_path, left=8800000, top=520000, width=2600000, height=807000)

    # Content slides (2-13): Compact watermark/logo in top right corner
    for idx in range(1, 13):
        slide = prs.slides[idx]
        slide.shapes.add_picture(logo_path, left=10350000, top=420000, width=1300000, height=403000)
    print("Sapienza logo added to all slides!")

# Save presentation
prs.save(pptx_path)
print("SUCCESS: Saved presentation to:", pptx_path)
