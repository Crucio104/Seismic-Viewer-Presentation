import pptx
from pptx.util import Pt
from pptx.dml.color import RGBColor

BORDEAUX      = RGBColor(130, 36, 51)    # #822433
SLATE_NAVY    = RGBColor(32, 64, 96)     # #204060
SAGE_TEAL     = RGBColor(35, 107, 94)    # #236B5E
TERRACOTTA    = RGBColor(184, 88, 40)    # #B85828
CHARCOAL_DARK = RGBColor(28, 24, 25)     # #1C1819

TINT_BORDEAUX = RGBColor(251, 244, 243)  # #FBF4F3
TINT_NAVY     = RGBColor(242, 246, 250)  # #F2F6FA
TINT_SAGE     = RGBColor(242, 248, 246)  # #F2F8F6
TINT_AMBER    = RGBColor(253, 248, 242)  # #FDF8F2

prs = pptx.Presentation('thesis/presentation/Seismic Viewer - Thesis Defense.pptx')

# -------------------------------------------------------------
# SLIDE 7: 4D Temporal Animation
# -------------------------------------------------------------
s7 = prs.slides[6]
for sh in s7.shapes:
    if sh.name == 'Google Shape;206;p19': # Card 1
        sh.top = 2150000; sh.height = 1750000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_AMBER
        sh.line.color.rgb = RGBColor(235, 215, 197)
    elif sh.name == 'Google Shape;207;p19': # Card 2
        sh.top = 2150000; sh.height = 1750000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_NAVY
        sh.line.color.rgb = RGBColor(208, 220, 229)
    elif sh.name == 'Google Shape;208;p19': # Card 3
        sh.top = 2150000; sh.height = 1750000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_SAGE
        sh.line.color.rgb = RGBColor(205, 224, 219)
    elif sh.name in ['Google Shape;211;p19', 'Google Shape;212;p19', 'Google Shape;213;p19']:
        sh.top = 2370000
    elif sh.name == 'Google Shape;221;p19': # Card 1 title
        sh.top = 2350000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16.5); p.font.bold = True; p.font.color.rgb = TERRACOTTA
            for r in p.runs: r.font.size = Pt(16.5); r.font.bold = True; r.font.color.rgb = TERRACOTTA
    elif sh.name == 'Google Shape;222;p19': # Card 2 title
        sh.top = 2350000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16.5); p.font.bold = True; p.font.color.rgb = SLATE_NAVY
            for r in p.runs: r.font.size = Pt(16.5); r.font.bold = True; r.font.color.rgb = SLATE_NAVY
    elif sh.name == 'Google Shape;223;p19': # Card 3 title
        sh.top = 2350000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16.5); p.font.bold = True; p.font.color.rgb = SAGE_TEAL
            for r in p.runs: r.font.size = Pt(16.5); r.font.bold = True; r.font.color.rgb = SAGE_TEAL
    elif sh.name in ['Google Shape;216;p19', 'Google Shape;217;p19', 'Google Shape;218;p19']:
        sh.top = 2750000; sh.height = 950000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13.5); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13.5); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;205;p19': # Bottom card
        sh.top = 4150000; sh.height = 1250000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_BORDEAUX
        sh.line.color.rgb = RGBColor(232, 213, 216)
    elif sh.name in ['Google Shape;209;p19', 'Google Shape;210;p19']:
        sh.top = 4400000
    elif sh.name == 'Google Shape;219;p19': # Bottom card title
        sh.top = 4320000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16.5); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(16.5); r.font.bold = True; r.font.color.rgb = BORDEAUX
    elif sh.name == 'Google Shape;220;p19': # Bottom card body
        sh.top = 4700000; sh.height = 550000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13.5); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13.5); r.font.color.rgb = CHARCOAL_DARK

# -------------------------------------------------------------
# SLIDE 9: Clustering Paradigms
# -------------------------------------------------------------
s9 = prs.slides[8]
for sh in s9.shapes:
    if sh.name == 'Google Shape;245;p21': # Card 1 (DBSCAN)
        sh.top = 2150000; sh.height = 1850000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_SAGE
        sh.line.color.rgb = RGBColor(205, 224, 219)
    elif sh.name == 'Google Shape;246;p21': # Card 2 (KMeans)
        sh.top = 2150000; sh.height = 1850000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_NAVY
        sh.line.color.rgb = RGBColor(208, 220, 229)
    elif sh.name == 'Google Shape;247;p21': # Card 3 (Cylinder Grid)
        sh.top = 2150000; sh.height = 1850000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_AMBER
        sh.line.color.rgb = RGBColor(235, 215, 197)
    elif sh.name in ['Google Shape;248;p21', 'Google Shape;249;p21', 'Google Shape;250;p21']:
        sh.top = 2370000
    elif sh.name == 'Google Shape;261;p21': # DBSCAN title
        sh.top = 2350000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = SAGE_TEAL
            for r in p.runs: r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = SAGE_TEAL
    elif sh.name == 'Google Shape;262;p21': # KMeans title
        sh.top = 2350000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = SLATE_NAVY
            for r in p.runs: r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = SLATE_NAVY
    elif sh.name == 'Google Shape;263;p21': # Cylinder Grid title
        sh.top = 2350000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = TERRACOTTA
            for r in p.runs: r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = TERRACOTTA
    elif sh.name in ['Google Shape;255;p21', 'Google Shape;257;p21', 'Google Shape;259;p21']:
        sh.top = 2750000; sh.height = 600000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name in ['Google Shape;256;p21', 'Google Shape;258;p21', 'Google Shape;260;p21']:
        sh.top = 3450000; sh.height = 450000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;244;p21': # Bottom card
        sh.top = 4250000; sh.height = 1150000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_BORDEAUX
        sh.line.color.rgb = RGBColor(232, 213, 216)
    elif sh.name == 'Google Shape;253;p21': # Bottom card title
        sh.top = 4380000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16.5); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(16.5); r.font.bold = True; r.font.color.rgb = BORDEAUX
    elif sh.name == 'Google Shape;254;p21': # Bottom card body
        sh.top = 4720000; sh.height = 550000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13.5); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13.5); r.font.color.rgb = CHARCOAL_DARK

prs.save('thesis/presentation/Seismic Viewer - Thesis Defense.pptx')
print('Slides 7 & 9 polished successfully!')
