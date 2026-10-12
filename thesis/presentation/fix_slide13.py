import pptx
from pptx.util import Pt
from pptx.dml.color import RGBColor

BORDEAUX      = RGBColor(130, 36, 51)    # #822433
SLATE_NAVY    = RGBColor(32, 64, 96)     # #204060
SAGE_TEAL     = RGBColor(35, 107, 94)    # #236B5E
TERRACOTTA    = RGBColor(184, 88, 40)    # #B85828
CHARCOAL_DARK = RGBColor(28, 24, 25)     # #1C1819
TEXT_MUTED    = RGBColor(90, 80, 82)     # #5A5052

TINT_BORDEAUX = RGBColor(251, 244, 243)  # #FBF4F3
TINT_NAVY     = RGBColor(242, 246, 250)  # #F2F6FA
TINT_SAGE     = RGBColor(242, 248, 246)  # #F2F8F6
TINT_AMBER    = RGBColor(253, 248, 242)  # #FDF8F2

prs = pptx.Presentation('thesis/presentation/Seismic Viewer - Thesis Defense.pptx')
s13 = prs.slides[12]

# 4 Cards Layout
# Top-Left: Monte Carlo (Shape 327)
# Bottom-Left: Real Catalogue (Shape 328)
# Top-Right: Serialization (Shape 329)
# Bottom-Right: Algorithmic Throughput (Shape 330)

for sh in s13.shapes:
    # Top-Left Card: Monte Carlo
    if sh.name == 'Google Shape;327;p25':
        sh.left = 609600; sh.top = 1950000; sh.width = 5353050; sh.height = 1900000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_SAGE
        sh.line.color.rgb = RGBColor(205, 224, 219)
    elif sh.name == 'Google Shape;333;p25': # Monte Carlo Title
        sh.left = 847725; sh.top = 2050000; sh.width = 4972050; sh.height = 250000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16.5); p.font.bold = True; p.font.color.rgb = SAGE_TEAL
            for r in p.runs: r.font.size = Pt(16.5); r.font.bold = True; r.font.color.rgb = SAGE_TEAL
    elif sh.name == 'Google Shape;334;p25': # Monte Carlo Body
        sh.left = 847725; sh.top = 2380000; sh.width = 4972050; sh.height = 650000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;340;p25': # < 0.51°
        sh.left = 847725; sh.top = 3100000; sh.width = 1600000; sh.height = 420000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(24); p.font.bold = True; p.font.color.rgb = SAGE_TEAL
            for r in p.runs: r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = SAGE_TEAL
    elif sh.name == 'Google Shape;345;p25': # Dip MAE
        sh.left = 847725; sh.top = 3520000; sh.width = 2000000; sh.height = 200000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = TEXT_MUTED
            for r in p.runs: r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = TEXT_MUTED
    elif sh.name == 'Google Shape;341;p25': # < 0.93°
        sh.left = 2800000; sh.top = 3100000; sh.width = 1600000; sh.height = 420000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(24); p.font.bold = True; p.font.color.rgb = SAGE_TEAL
            for r in p.runs: r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = SAGE_TEAL
    elif sh.name == 'Google Shape;346;p25': # Strike MAE
        sh.left = 2800000; sh.top = 3520000; sh.width = 2000000; sh.height = 200000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = TEXT_MUTED
            for r in p.runs: r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = TEXT_MUTED

    # Bottom-Left Card: Real Catalogue
    elif sh.name == 'Google Shape;328;p25':
        sh.left = 609600; sh.top = 4050000; sh.width = 5353050; sh.height = 1450000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_BORDEAUX
        sh.line.color.rgb = RGBColor(232, 213, 216)
    elif sh.name == 'Google Shape;335;p25': # Title
        sh.left = 847725; sh.top = 4180000; sh.width = 4972050; sh.height = 250000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16.5); p.font.bold = True; p.font.color.rgb = BORDEAUX
            for r in p.runs: r.font.size = Pt(16.5); r.font.bold = True; r.font.color.rgb = BORDEAUX
    elif sh.name == 'Google Shape;336;p25': # Body
        sh.left = 847725; sh.top = 4500000; sh.width = 4972050; sh.height = 850000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK

    # Top-Right Card: Serialization
    elif sh.name == 'Google Shape;329;p25':
        sh.left = 6229350; sh.top = 1950000; sh.width = 5353050; sh.height = 1900000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_AMBER
        sh.line.color.rgb = RGBColor(235, 215, 197)
    elif sh.name == 'Google Shape;337;p25': # Title
        sh.left = 6467475; sh.top = 2050000; sh.width = 4972050; sh.height = 250000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16.5); p.font.bold = True; p.font.color.rgb = TERRACOTTA
            for r in p.runs: r.font.size = Pt(16.5); r.font.bold = True; r.font.color.rgb = TERRACOTTA
    elif sh.name == 'Google Shape;342;p25': # Bullet 1
        sh.left = 6467475; sh.top = 2450000; sh.width = 4972050; sh.height = 300000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK
    elif sh.name == 'Google Shape;343;p25': # Bullet 2
        sh.left = 6467475; sh.top = 2850000; sh.width = 4972050; sh.height = 300000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = TERRACOTTA
            for r in p.runs: r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = TERRACOTTA
    elif sh.name == 'Google Shape;344;p25': # Bullet 3
        sh.left = 6467475; sh.top = 3250000; sh.width = 4972050; sh.height = 300000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK

    # Bottom-Right Card: Throughput
    elif sh.name == 'Google Shape;330;p25':
        sh.left = 6229350; sh.top = 4050000; sh.width = 5353050; sh.height = 1450000
        sh.fill.solid(); sh.fill.fore_color.rgb = TINT_NAVY
        sh.line.color.rgb = RGBColor(208, 220, 229)
    elif sh.name == 'Google Shape;338;p25': # Title
        sh.left = 6467475; sh.top = 4180000; sh.width = 4972050; sh.height = 250000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(16.5); p.font.bold = True; p.font.color.rgb = SLATE_NAVY
            for r in p.runs: r.font.size = Pt(16.5); r.font.bold = True; r.font.color.rgb = SLATE_NAVY
    elif sh.name == 'Google Shape;339;p25': # Body
        sh.left = 6467475; sh.top = 4500000; sh.width = 4972050; sh.height = 850000
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(13); p.font.color.rgb = CHARCOAL_DARK
            for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = CHARCOAL_DARK

prs.save('thesis/presentation/Seismic Viewer - Thesis Defense.pptx')
print('Slide 13 fixed successfully!')
